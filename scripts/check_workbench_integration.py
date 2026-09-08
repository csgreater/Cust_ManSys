"""Synthetic, authenticated integration checks in an exclusively created local DB.

Run with .venv/Scripts/python scripts/check_workbench_integration.py.
The test database is retained for browser QA; its name contains no credentials.
--serve NAME serves that test database on localhost:8918 without creating data.
--drop NAME deletes only a database recorded as created by this script.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import uuid
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from app import db
from app.config import settings

OUTPUT = ROOT / "output/implementation-20260908"
PREFIX = "codex_workbench_test_"


def select_database(name):
    cfg = db.parse_database_url()
    if cfg.host not in {"127.0.0.1", "localhost", "::1"} or settings.is_production:
        raise RuntimeError("Integration fixtures require loopback development MySQL")
    if not re.fullmatch(PREFIX + r"[a-f0-9]{12}", name):
        raise ValueError("Not a disposable workbench test database")
    parts = urlsplit(settings.database_url)
    db.settings = replace(settings, database_url=urlunsplit(parts._replace(path="/" + name)))


def create_fixtures(name):
    select_database(name)
    # Deliberately no IF NOT EXISTS: an existing database must never be reused.
    with db.connection(database="") as conn:
        with conn.cursor() as cur:
            cur.execute(f"CREATE DATABASE `{name}` CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci")
    OUTPUT.mkdir(parents=True, exist_ok=True)
    (OUTPUT / f"{name}.json").write_text(json.dumps({"database": name, "synthetic": True}), encoding="utf-8")
    from scripts.init_db import create_tables, seed_data
    from scripts.migrate_anomaly_workbench import migrate
    from app.analytics_aggregates import create_dashboard_tables, create_product_monthly_table, refresh_dashboard_months, refresh_product_months
    create_tables()
    with db.connection() as conn:
        create_dashboard_tables(conn)
        create_product_monthly_table(conn)
        migrate(conn)
        migrate(conn)  # idempotence
    seed_data()
    # Hand-derived revenue/profit: June 5500/380, May 4200/700.
    fixtures = [
        ("P1", "06", 1000, -100, 300), ("P2", "06", 2000, 80, 100),
        ("P3", "06", 1000, 200, 250), ("P4", "06", 1500, 200, 0),
        ("P1", "05", 1200, 100, 100), ("P2", "05", 1000, 100, 50),
        ("P3", "05", 1000, 200, 50), ("P5", "05", 1000, 300, 0),
    ]
    with db.connection() as conn:
        with conn.cursor() as cur:
            for product, month, revenue, profit, fees in fixtures:
                cur.execute("""INSERT INTO t_order_sku_detail
                  (customer_no,dept,platform,shop_name,order_no,sku_id,category,product_classification,
                   product_name,product_no,unit,qty,share_receivable,cost,freight,ship_time,province,city,order_source)
                  VALUES ('TEST','华东','平台A','门店A',%s,%s,'食品','常温',%s,%s,'件',10,%s,%s,%s,%s,'上海','上海','直营')""",
                            (f"{month}-{product}", f"SKU-{product}", f"测试商品 {product}", product, revenue, revenue-profit-fees, fees, f"2026-{month}-15"))
            cur.execute("""INSERT INTO t_order_sku_detail
              (customer_no,dept,platform,shop_name,order_no,category,product_name,product_no,unit,share_receivable,cost,ship_time)
              VALUES ('TEST','华南','平台B','门店B','HIDDEN','食品','隔离商品','HIDDEN','件',9000,10000,'2026-06-15')""")
            cur.execute("UPDATE t_order_sku_detail SET share_receivable=900 WHERE order_no='06-P1'")
            cur.execute("""INSERT INTO t_order_sku_detail
              (customer_no,dept,platform,shop_name,order_no,sku_id,category,product_classification,
               product_name,product_no,unit,share_receivable,ship_time,province,city,order_source)
              VALUES ('TEST','华东','平台A','门店A','06-P1','SKU-P1','食品','常温',
                      '同货号别名','P1','件',100,'2026-06-15','上海','上海','直营')""")
            for owner, batch, shop, dept, platform, count in [
                ('admin','QA-ERRORS','门店A','华东','平台A',250),
                ('importer','OWN-ERRORS','门店A','华东','平台A',1),
                ('admin','HIDDEN-ERROR','门店B','华南','平台B',1),
            ]:
                cur.execute("""INSERT INTO t_import_log (batch_no,import_user,file_name,total_rows,fail_rows,status)
                  VALUES (%s,%s,'synthetic.xlsx',%s,%s,'failed')""", (batch,owner,count,count))
                cur.executemany("""INSERT INTO tmp_order_import
                  (batch_no,row_no,dept,platform,shop_name,product_no,product_name,error_message)
                  VALUES (%s,%s,%s,%s,%s,'BAD','待修复测试商品','发货日期无效')""",
                                [(batch,i+2,dept,platform,shop) for i in range(count)])
        refresh_product_months(conn, ['2026-05-01','2026-06-01'])
        refresh_dashboard_months(conn, ['2026-05-01','2026-06-01'])


def verify():
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from app.main import app
    results = []
    filters = {"start_time":"2026-06-01","end_time":"2026-06-30","shop_name":"门店A"}

    def checked(response, code=200):
        assert response.status_code == code, f"{response.request.url.path}: {response.status_code} {response.text[:300]}"
        return response.json().get("data") if code == 200 and 'json' in response.headers.get('content-type','') else response

    def login(client, username):
        checked(client.post('/api/login', data={'username':username,'password':getattr(settings,username+'_password')}))

    with TestClient(app) as client, TestClient(app) as other:
        checked(client.get('/api/exceptions',params=filters),401)
        login(client,'admin')
        ctx = checked(client.get('/api/workbench/context',params=filters))
        assert ctx['latest_period'] == {'start_time':'2026-06-01','end_time':'2026-06-30'}
        operating = checked(client.get('/api/exceptions',params=filters))
        aggregate=checked(client.get('/api/analytics/products',params=filters))['rows']
        details=checked(client.get('/api/analytics/products',params={**filters,'province':'上海'}))['rows']
        assert aggregate==details, 'Product aliases split only in detail fallback'
        assert {'negative_profit','low_margin','high_fee','revenue_decline','profit_decline'} <= {r['rule_id'] for r in operating['rows']}
        assert all(r['shop_name']=='门店A' for r in operating['rows'])
        row = operating['rows'][0]
        review = {'filters':filters,'id':row['id'],'evidence_hash':row['evidence_hash'],'status':'confirmed','note':'已检查合成证据'}
        checked(client.post('/api/exceptions/review',json=review))
        confirmed = checked(client.get('/api/exceptions',params={**filters,'status':'confirmed'}))
        assert confirmed['total']==1
        checked(client.post('/api/exceptions/review',json={**review,'evidence_hash':'old'}),409)
        results.append('scoped operating exceptions, persisted review and stale evidence refusal')
        quality = checked(client.get('/api/exceptions',params={**filters,'type':'quality','batch_no':'QA-ERRORS','page':5,'page_size':50}))
        assert quality['total']==250 and len(quality['rows'])==50
        assert len({r['id'] for r in quality['rows']})==50
        export = checked(client.get('/api/imports/QA-ERRORS/errors/export.xlsx'))
        wb = load_workbook(BytesIO(export.content),data_only=False)
        assert wb.active.max_row==251, wb.active.max_row
        (OUTPUT/'quality-250.xlsx').write_bytes(export.content)
        results.append('250 invalid-date quality rows, full pagination and XLSX export')
        # Exercise the actual upload/worker/commit path with a July-only fixture.
        from openpyxl import Workbook
        from app.import_service import HEADER_MAP
        fixture={'customer_no':'TEST','dept':'华东','platform':'平台A','shop_name':'门店A',
                 'order_no':'UPLOAD-ONLY','product_no':'UPLOAD','product_name':'上传验证商品',
                 'share_receivable':100,'cost':70,'excel_profit':30,'ship_time':'2026-07-15','qty':1}
        workbook=Workbook(); workbook.active.append(list(HEADER_MAP))
        workbook.active.append([fixture.get(field,'') for field in HEADER_MAP.values()])
        content=BytesIO(); workbook.save(content)
        upload=checked(client.post('/api/imports/upload',files={'file':('integration.xlsx',content.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}))
        batch=upload['batch_no']
        assert checked(client.get('/api/imports/'+batch))['log']['status']=='validated'
        checked(client.post('/api/imports/'+batch+'/commit'))
        checked(client.post('/api/imports/'+batch+'/commit'))
        with db.connection() as conn:
            with conn.cursor() as cur:
                cur.execute('SELECT COUNT(*) AS n,MIN(source_row_no) AS row_no FROM t_order_sku_detail WHERE source_batch_no=%s',(batch,))
                assert cur.fetchone()=={'n':1,'row_no':2}
                cur.execute("""INSERT INTO t_import_log(batch_no,import_user,status,import_time)
                  VALUES ('ABANDONED','admin','processing',NOW()-INTERVAL 20 MINUTE)""")
        from app.import_lifecycle import processing_lease
        with db.connection() as worker:
            with processing_lease(worker,'ABANDONED'):
                checked(client.post('/api/imports/ABANDONED/recover'),409)
        checked(client.post('/api/imports/ABANDONED/recover'))
        assert checked(client.get('/api/imports/ABANDONED'))['log']['status']=='failed'
        results.append('actual XLSX upload, atomic idempotent commit, source lineage and abandoned-worker recovery')
        report = checked(client.post('/api/reports',json={'filters':filters,'title':'合成数据经营报告','report_type':'monthly'}))
        assert float(report['summary']['revenue'])==5500 and float(report['summary']['profit'])==380
        assert float(report['previous_summary']['revenue'])==4200
        assert sum(float(r['revenue_change']) for r in report['contributions'])==1300
        saved = checked(client.get('/api/reports/'+report['id']))
        assert saved['stale'] is False
        newer = checked(client.post('/api/reports/'+report['id']+'/regenerate'))
        assert newer['id'] != report['id'] and newer['version']==report['version']+1
        assert checked(client.get('/api/reports/'+report['id']))['version']==report['version']
        from concurrent.futures import ThreadPoolExecutor
        from app.workbench_store import save_report
        from app.permissions import load_user_context
        with db.connection() as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT id FROM t_user WHERE username='admin'")
                user=load_user_context(conn,cur.fetchone()['id'])
        def concurrent_save(_):
            with db.connection() as conn:
                return save_report(conn,user,report['body'],['analytics','import'])['version']
        with ThreadPoolExecutor(max_workers=2) as pool:
            versions=list(pool.map(concurrent_save,range(2)))
        assert sorted(versions)==[3,4]
        for ext in ['xlsx','pdf']:
            downloaded=checked(client.get('/api/reports/'+newer['id']+'/export.'+ext))
            assert downloaded.content.startswith(b'%PDF' if ext=='pdf' else b'PK')
            (OUTPUT/('sample-report.'+ext)).write_bytes(downloaded.content)
        results.append('hand-derived totals, contribution reconciliation, immutable and concurrent versions, PDF/XLSX')
        checked(client.post('/api/exceptions/review',json={**review,'note':'已复核，更新核查备注'}))
        assert checked(client.get('/api/reports/'+report['id']))['stale'] is True
        results.append('review changes mark archived report stale without rewriting its body')
        for params in [{'start_time':'2026-07-01','end_time':'2026-06-01'},{**filters,'page_size':201}]:
            checked(client.get('/api/exceptions',params=params),400)
        checked(client.put('/api/workbench/rules',json={'rules':{'low_margin_pct':-1}}),400)
        login(other,'importer')
        checked(other.get('/api/imports/QA-ERRORS'),404)
        own=checked(other.get('/api/exceptions',params={**filters,'type':'quality'}))
        assert own['total']==1
        checked(other.get('/api/exceptions',params=filters),403)
        checked(other.get('/api/reports/'+report['id']),403)
        with db.connection() as conn:
            with conn.cursor() as cur:
                cur.execute("UPDATE t_role SET permissions='view,import,analytics,export' WHERE role_code='importer'")
        quality_report=checked(other.post('/api/reports',json={'filters':filters}))
        assert len(quality_report['quality_rows'])==1
        with db.connection() as conn:
            with conn.cursor() as cur:
                cur.execute("""UPDATE t_role_data_scope s JOIN t_role r ON r.id=s.role_id
                  SET s.scope_value='门店A' WHERE r.role_code='importer' AND s.scope_type='shop'""")
        checked(other.get('/api/reports/'+quality_report['id']))
        results.append('quality report retains explicit shop scope when unrelated access is revoked')
        login(other,'analyst')
        checked(other.get('/api/reports/'+report['id']),404)
        own_report=checked(other.post('/api/reports',json={'filters':{k:v for k,v in filters.items() if k!='shop_name'}}))
        assert own_report['quality_rows']==[]
        with db.connection() as conn:
            with conn.cursor() as cur:
                cur.execute("""UPDATE t_role_data_scope s JOIN t_role r ON r.id=s.role_id
                  SET s.scope_value='门店A' WHERE r.role_code='analyst' AND s.scope_type='shop'""")
        checked(other.get('/api/reports/'+own_report['id']),403)
        visible=checked(other.get('/api/reports'))
        assert not visible['reports']
        scoped=checked(other.get('/api/exceptions',params={k:v for k,v in filters.items() if k!='shop_name'}))
        assert all(r['shop_name']=='门店A' for r in scoped['rows'])
        results.append('cross-owner denial, import permission isolation and narrowed historical report scope')
        # Rule version consistency and stale report signal.
        rules=ctx['rules']; rules['low_margin_pct']=6
        updated=checked(client.put('/api/workbench/rules',json={'rules':rules}))
        current=checked(client.get('/api/exceptions',params=filters))
        assert updated['version']==current['rules_version']
        assert checked(client.get('/api/reports/'+report['id']))['stale'] is True
        results.append('unified rule version and stale report detection')
        with db.connection() as conn:
            with conn.cursor() as cur:
                cur.execute("UPDATE t_user_role ur JOIN t_user u ON u.id=ur.user_id SET ur.role_id=(SELECT id FROM t_role WHERE role_code='importer') WHERE u.username='admin'")
        try:
            checked(client.get('/api/reports/'+report['id']),403)
            checked(client.get('/api/reports/'+report['id']+'/export.xlsx'),403)
            results.append('admin demotion blocks historical cross-owner import evidence')
        finally:
            with db.connection() as conn:
                with conn.cursor() as cur:
                    cur.execute("UPDATE t_user_role ur JOIN t_user u ON u.id=ur.user_id SET ur.role_id=(SELECT id FROM t_role WHERE role_code='admin') WHERE u.username='admin'")
    return results


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--serve')
    parser.add_argument('--drop')
    args=parser.parse_args()
    if args.serve or args.drop:
        name=args.serve or args.drop
        select_database(name)
        marker=OUTPUT/f'{name}.json'
        if not marker.exists() or json.loads(marker.read_text()) != {'database':name,'synthetic':True}:
            raise RuntimeError('Missing exclusive test database ownership marker')
        if args.drop:
            with db.connection(database='') as conn:
                with conn.cursor() as cur:
                    cur.execute(f'DROP DATABASE `{name}`')
            marker.unlink()
            print('Removed exclusively created synthetic test database.')
        else:
            import uvicorn
            from app.main import app
            uvicorn.run(app,host='127.0.0.1',port=8918,log_level='warning')
        return
    name=PREFIX+uuid.uuid4().hex[:12]
    create_fixtures(name)
    print('Synthetic database:',name,flush=True)
    results=verify()
    evidence={'database':name,'checks':results,'passed':True,'date':'2026-09-08'}
    (OUTPUT/'integration-evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(evidence,ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()
