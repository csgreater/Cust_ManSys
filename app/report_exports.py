"""XLSX and PDF exporters for deterministic anomaly reports."""
from __future__ import annotations

import io
import json
import math
import os
import re
from datetime import date, datetime
from decimal import Decimal
from html import escape
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
PHONE_RE = re.compile(r"(?<!\d)(1\d{2})\d{4}(\d{4})(?!\d)")
EMAIL_RE = re.compile(r"(?<![\w.])([\w.+-])[^@\s]*(@[^\s]+)")

EXCEPTION_COLUMNS = (
    ("id", "异常ID"),
    ("type", "类型"),
    ("rule_label", "规则"),
    ("severity", "严重程度"),
    ("object_label", "对象"),
    ("platform", "平台"),
    ("shop_name", "店铺"),
    ("product_no", "货号"),
    ("current_value", "当前值"),
    ("baseline_value", "基线值"),
    ("change_pct", "变化率(%)"),
    ("impact_amount", "规则影响额"),
    ("message", "触发说明"),
    ("suggestion", "核查建议"),
    ("status", "核查状态"),
    ("note", "核查备注"),
    ("reviewed_at", "核查时间"),
    ("evidence_hash", "证据哈希"),
    ("evidence", "证据"),
)

FILTER_LABELS = (
    ("start_time", "开始日期"),
    ("end_time", "结束日期"),
    ("comparison_mode", "比较方式"),
    ("dept", "部门"),
    ("platform", "平台"),
    ("shop_name", "店铺"),
    ("category", "产品大类"),
    ("product_classification", "货品分类"),
    ("product", "产品"),
    ("order_no", "订单号"),
    ("province", "省份"),
    ("city", "城市"),
    ("order_source", "订单来源"),
    ("batch_no", "导入批次"),
)


def _formula_safe(value: Any) -> Any:
    if isinstance(value, str):
        return f"'{value}" if value.startswith(FORMULA_PREFIXES) else value
    if isinstance(value, dict):
        return {str(key): _formula_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_formula_safe(item) for item in value]
    return value


def _mask_pii(value: str) -> str:
    value = PHONE_RE.sub(r"\1****\2", value)
    return EMAIL_RE.sub(r"\1***\2", value)


def _cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        safe = _formula_safe(value)
        return _mask_pii(json.dumps(safe, ensure_ascii=False, sort_keys=True, default=str))
    if isinstance(value, str):
        return _formula_safe(_mask_pii(value))
    if isinstance(value, (Decimal, int, float, date, datetime)):
        return value
    return _formula_safe(_mask_pii(str(value)))


def _numeric_cell(value: Any) -> Any:
    """Restore JSON-archived Decimal strings as real spreadsheet numbers."""
    if value in (None, ""):
        return ""
    if isinstance(value, bool):
        return value
    try:
        number = value if isinstance(value, Decimal) else Decimal(str(value))
    except (ArithmeticError, TypeError, ValueError):
        return _cell_value(value)
    if not number.is_finite():
        return ""
    return number


def _report_metadata_rows(report: dict[str, Any]) -> list[tuple[str, Any]]:
    rows: list[tuple[str, Any]] = [
        ("报告ID", report.get("id") or "未保存"),
        ("版本", report.get("version") if report.get("version") is not None else "未保存"),
        ("生成时间", report.get("created_at") or "未提供"),
        ("报告类型", {"monthly": "月度经营报告", "exception": "异常专题报告"}.get(report.get("report_type"), report.get("report_type") or "")),
    ]
    filters = report.get("filters") or {}
    for key, label in FILTER_LABELS:
        value = filters.get(key)
        if key == "comparison_mode":
            value = {"previous_period": "前一周期", "year_over_year": "去年同期"}.get(value, value)
        if value not in (None, ""):
            rows.append((label, value))
    comparison = report.get("comparison") or {}
    if comparison.get("label"):
        rows.append(("对比期间", comparison["label"]))
    rules = report.get("rules") or {}
    for key, label in (
        ("low_margin_pct", "低利润率阈值"),
        ("high_fee_pct", "扣减费用率阈值"),
        ("decline_pct", "下降阈值"),
    ):
        if rules.get(key) is not None:
            rows.append((label, f"{rules[key]}%"))
    if rules.get("min_revenue") is not None:
        rows.append(("最低销售额", rules["min_revenue"]))
    rows.extend([
        ("规则版本", report.get("rules_version") or "未提供"),
        ("数据版本", report.get("data_version") or "未提供"),
    ])
    return rows


def _pdf_operating_values(row: dict[str, Any]) -> dict[str, str]:
    return {
        "rule_label": _mask_pii(str(row.get("rule_label") or "")),
        "object_label": _mask_pii(str(row.get("object_label") or "")),
        "status": _mask_pii(str(row.get("status") or "pending")),
        "note": _mask_pii(str(row.get("note") or "")),
    }


def _style_sheet(sheet, *, freeze: str = "A2", auto_filter: bool = True) -> None:
    sheet.freeze_panes = freeze
    if auto_filter and sheet.max_row >= 1 and sheet.max_column >= 1:
        sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(48, max(10, max((len(str(cell.value or "")) for cell in column), default=10) + 2))
        sheet.column_dimensions[letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _append_exception_sheet(workbook: Workbook, title: str, rows: Iterable[dict[str, Any]]):
    sheet = workbook.create_sheet(title)
    sheet.append([label for _, label in EXCEPTION_COLUMNS])
    for row in rows:
        sheet.append([
            _numeric_cell(row.get(key)) if key in {"baseline_value", "change_pct", "impact_amount"}
            or (key == "current_value" and row.get("type") != "quality")
            else _cell_value(row.get(key))
            for key, _ in EXCEPTION_COLUMNS
        ])
    _style_sheet(sheet)
    return sheet


def _workbook_bytes(workbook: Workbook) -> bytes:
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def export_exceptions_xlsx(rows: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _append_exception_sheet(workbook, "异常清单", rows)
    workbook.properties.title = "异常完整清单"
    return _workbook_bytes(workbook)


def export_report_xlsx(report: dict[str, Any]) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "经营摘要"
    summary_sheet.append(["指标", "本期", "对比期"])
    labels = (
        ("revenue", "销售额"),
        ("profit", "利润"),
        ("profit_rate", "利润率(%)"),
        ("fees", "扣减费用"),
        ("fee_rate", "扣减费用率(%)"),
        ("qty", "销量"),
        ("orders", "订单数"),
        ("detail_rows", "明细行数"),
    )
    current = report.get("summary") or {}
    previous = report.get("previous_summary") or {}
    for key, label in labels:
        summary_sheet.append([label, _numeric_cell(current.get(key)), _numeric_cell(previous.get(key))])
    summary_sheet.append(["报告标题", _cell_value(report.get("title")), ""])
    for label, value in _report_metadata_rows(report):
        summary_sheet.append([
            label,
            _numeric_cell(value) if label in {"版本", "最低销售额"} else _cell_value(value),
            "",
        ])
    _style_sheet(summary_sheet, auto_filter=False)

    contribution_sheet = workbook.create_sheet("变化贡献")
    contribution_columns = (
        ("dept", "部门"), ("platform", "平台"), ("shop_name", "店铺"),
        ("product_no", "货号"), ("category", "产品大类"),
        ("product_classification", "货品分类"), ("object_label", "对象"),
        ("current_revenue", "本期销售额"), ("previous_revenue", "对比期销售额"),
        ("revenue_change", "销售额贡献"), ("current_profit", "本期利润"),
        ("previous_profit", "对比期利润"), ("profit_change", "利润贡献"),
        ("disappeared", "本期消失"), ("new_in_period", "本期新增"),
    )
    contribution_sheet.append([label for _, label in contribution_columns])
    for row in report.get("contributions") or []:
        contribution_sheet.append([
            _numeric_cell(row.get(key))
            if key in {
                "current_revenue", "previous_revenue", "revenue_change",
                "current_profit", "previous_profit", "profit_change",
            }
            else _cell_value(row.get(key))
            for key, _ in contribution_columns
        ])
    _style_sheet(contribution_sheet)

    _append_exception_sheet(workbook, "经营异常", report.get("operating_rows") or [])
    _append_exception_sheet(workbook, "数据质量异常", report.get("quality_rows") or [])

    notes = workbook.create_sheet("口径与说明")
    notes.append(["项目", "内容"])
    notes.append(["提示", "报告中的变化贡献只表示数值差额，不代表已验证的业务原因。"])
    for warning in report.get("warnings") or []:
        notes.append(["口径提示", _cell_value(warning)])
    for line in report.get("narrative") or []:
        notes.append(["模板说明", _cell_value(line)])
    notes.append(["规则", _cell_value(report.get("rules") or {})])
    _style_sheet(notes, auto_filter=False)

    workbook.properties.title = str(report.get("title") or "分析报告")
    return _workbook_bytes(workbook)


def _pdf_font_name() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont, TTFError

    name = "CusManChinese"
    if name in pdfmetrics.getRegisteredFontNames():
        return name
    candidates = (
        r"C:\Windows\Fonts\simhei.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
    )
    for path in candidates:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont(name, path))
                return name
            except (TTFError, ValueError, OSError):
                # Some Noto TTC collections use unsupported CFF outlines.
                # Continue to another CJK face or the built-in CID fallback.
                continue
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light"


def _pdf_text(value: Any) -> str:
    return escape(_mask_pii(str(value if value is not None else "")))


def _pdf_number(value: Any) -> str:
    if value in (None, ""):
        return "—"
    if isinstance(value, Decimal):
        return format(value.quantize(Decimal("0.01")), ",f") if value.is_finite() else "—"
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return "—"
        return f"{value:,.2f}" if not isinstance(value, int) else f"{value:,}"
    if isinstance(value, str):
        try:
            number = Decimal(value)
        except ArithmeticError:
            return _pdf_text(value)
        return format(number.quantize(Decimal("0.01")), ",f") if number.is_finite() else "—"
    return _pdf_text(value)


def export_report_pdf(report: dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font_name = _pdf_font_name()
    stream = io.BytesIO()
    document = SimpleDocTemplate(
        stream,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        pageCompression=0,
        title=str(report.get("title") or "分析报告"),
        subject="Top 10 summaries and complete-data metrics",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ChineseTitle", parent=styles["Title"], fontName=font_name,
        fontSize=20, leading=26, alignment=TA_CENTER, textColor=colors.HexColor("#17365D"),
    )
    heading_style = ParagraphStyle(
        "ChineseHeading", parent=styles["Heading2"], fontName=font_name,
        fontSize=13, leading=18, textColor=colors.HexColor("#1F4E78"), spaceBefore=8, spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "ChineseBody", parent=styles["BodyText"], fontName=font_name,
        fontSize=9, leading=14, wordWrap="CJK",
    )
    small_style = ParagraphStyle(
        "ChineseSmall", parent=body_style, fontSize=7, leading=9, splitLongWords=True,
    )

    story: list[Any] = [Paragraph(_pdf_text(report.get("title") or "分析报告"), title_style), Spacer(1, 6 * mm)]
    filters = report.get("filters") or {}
    comparison = report.get("comparison") or {}
    story.append(Paragraph(
        f"报告期间：{_pdf_text(filters.get('start_time'))} 至 {_pdf_text(filters.get('end_time'))}　"
        f"对比期间：{_pdf_text(comparison.get('label') or '无有效基线')}",
        body_style,
    ))
    story.append(Paragraph("报告口径", heading_style))
    metadata_data: list[list[Any]] = [["项目", "内容"]]
    for label, value in _report_metadata_rows(report):
        metadata_data.append([
            Paragraph(_pdf_text(label), small_style),
            Paragraph(_pdf_text(value), small_style),
        ])
    metadata_table = Table(metadata_data, colWidths=[35 * mm, 145 * mm], repeatRows=1)
    metadata_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B4C6E7")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.extend([metadata_table, Spacer(1, 3 * mm)])
    story.append(Paragraph("经营摘要", heading_style))
    summary = report.get("summary") or {}
    previous = report.get("previous_summary") or {}
    metric_rows = [["指标", "本期", "对比期"]]
    for key, label in (
        ("revenue", "销售额"), ("profit", "利润"), ("profit_rate", "利润率(%)"),
        ("fees", "扣减费用"), ("qty", "销量"), ("orders", "订单数"),
    ):
        metric_rows.append([label, _pdf_number(summary.get(key)), _pdf_number(previous.get(key))])
    metrics = Table(metric_rows, colWidths=[48 * mm, 55 * mm, 55 * mm], repeatRows=1)
    metrics.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B4C6E7")),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6FA")]),
    ]))
    story.extend([metrics, Spacer(1, 4 * mm)])
    for line in report.get("narrative") or []:
        story.append(Paragraph(f"- {_pdf_text(line)}", body_style))
    for warning in report.get("warnings") or []:
        story.append(Paragraph(f"口径提示：{_pdf_text(warning)}", body_style))

    story.append(PageBreak())
    top_n = 10
    operating = list(report.get("operating_rows") or [])
    story.append(Paragraph(f"重点经营异常（Top {top_n}）", heading_style))
    story.append(Paragraph(f"本页展示 Top {min(top_n, len(operating))} / 共 {len(operating)} 条；完整清单请使用 Excel 导出。", body_style))
    op_rows: list[list[Any]] = [["规则", "对象", "当前值", "基线", "影响额", "核查状态", "核查备注"]]
    review_notes: list[tuple[str, str]] = []
    for row in operating[:top_n]:
        review = _pdf_operating_values(row)
        if review["note"]:
            review_notes.append((review["object_label"], review["note"]))
        op_rows.append([
            Paragraph(_pdf_text(review["rule_label"]), small_style),
            Paragraph(_pdf_text(review["object_label"]), small_style),
            _pdf_number(row.get("current_value")),
            _pdf_number(row.get("baseline_value")),
            _pdf_number(row.get("impact_amount")),
            Paragraph(_pdf_text({"pending": "待核查", "confirmed": "已确认问题", "dismissed": "已排除"}.get(review["status"], review["status"])), small_style),
            Paragraph("见表后备注" if review["note"] else "—", small_style),
        ])
    if len(op_rows) == 1:
        op_rows.append(["—", "暂无经营异常", "—", "—", "—", "—", "—"])
    op_table = Table(op_rows, colWidths=[19 * mm, 32 * mm, 18 * mm, 18 * mm, 18 * mm, 20 * mm, 55 * mm], repeatRows=1)
    op_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B4C6E7")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(op_table)
    for index, (object_label, note) in enumerate(review_notes, start=1):
        story.append(Paragraph(f"核查备注 {index}｜{_pdf_text(object_label)}：{_pdf_text(note)}", small_style))

    quality = list(report.get("quality_rows") or [])
    story.append(Paragraph(f"数据质量异常（Top {top_n}）", heading_style))
    story.append(Paragraph("质量异常不按发货日期过滤，仍服从其他筛选与批次权限。", small_style))
    story.append(Paragraph(f"本页展示 Top {min(top_n, len(quality))} / 共 {len(quality)} 条；完整清单请使用 Excel 导出。", body_style))
    quality_rows: list[list[Any]] = [["批次与行号", "级别", "校验信息", "建议"]]
    for row in quality[:top_n]:
        quality_rows.append([
            Paragraph(_pdf_text(row.get("object_label")), small_style),
            _pdf_text({"high": "高", "medium": "中", "low": "低"}.get(row.get("severity"), row.get("severity"))),
            Paragraph(_pdf_text(row.get("message")), small_style),
            Paragraph(_pdf_text(row.get("suggestion")), small_style),
        ])
    if len(quality_rows) == 1:
        quality_rows.append(["—", "—", "暂无数据质量异常", "—"])
    quality_table = Table(quality_rows, colWidths=[42 * mm, 20 * mm, 65 * mm, 53 * mm], repeatRows=1)
    quality_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B4C6E7")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(quality_table)

    story.append(Paragraph(f"变化贡献（Top {top_n}）", heading_style))
    count = len(report.get("contributions") or [])
    story.append(Paragraph(f"展示 {min(top_n, count)} / 共 {count} 个对象。", small_style))
    story.append(Paragraph("变化贡献只表示数值差额，不代表已验证的业务原因。", body_style))
    contribution_rows: list[list[Any]] = [["对象", "销售额贡献", "利润贡献"]]
    for row in (report.get("contributions") or [])[:top_n]:
        contribution_rows.append([
            Paragraph(_pdf_text(row.get("object_label")), small_style),
            _pdf_number(row.get("revenue_change")),
            _pdf_number(row.get("profit_change")),
        ])
    if len(contribution_rows) == 1:
        contribution_rows.append(["暂无贡献数据", "—", "—"])
    contribution_table = Table(contribution_rows, colWidths=[100 * mm, 40 * mm, 40 * mm], repeatRows=1)
    contribution_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B4C6E7")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(contribution_table)

    def page_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(font_name, 8)
        canvas.drawRightString(A4[0] - 15 * mm, 8 * mm, f"第 {doc.page} 页")
        canvas.setSubject("Top 10 summaries and complete-data metrics")
        canvas.restoreState()

    document.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    return stream.getvalue()
