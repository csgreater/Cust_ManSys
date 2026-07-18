from __future__ import annotations

import csv
import gc
import hashlib
import hmac
from io import BytesIO
import json
import logging
import tempfile
import time
import asyncio
from collections import defaultdict, deque
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, Literal

from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.encoders import jsonable_encoder
from fastapi.responses import HTMLResponse, JSONResponse, PlainTextResponse, RedirectResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel
from pymysql.cursors import SSDictCursor
from starlette.middleware.sessions import SessionMiddleware
from openpyxl import Workbook

from app.config import BASE_DIR, settings
from app.analytics_aggregates import (
    can_use_global_monthly,
    can_use_product_monthly,
    product_monthly_where,
    refresh_dashboard_months,
    refresh_product_months,
)
from app.ark_analytics import ark_analytics_enabled, parse_analytics_intent
from app.ark_coding import ArkCodingError, ask_coding_plan, coding_chat_enabled
from app.db import connection
from app.import_service import HEADER_MAP, ImportParseResult, iter_excel_rows, new_batch_no, parse_excel
from app.nl_analytics import audit_sql_plan, build_chart, build_sql, normalize_rows, parse_question_filters, summarize_answer
from app.permissions import has_permission, load_user_context, requested_scope_filters, scope_clause
from app.security import verify_password


settings.validate_for_runtime()

app = FastAPI(title="订单数据管理系统")
app.add_middleware(
    SessionMiddleware,
    secret_key=settings.secret_key,
    same_site=settings.session_same_site,
    https_only=settings.secure_cookies,
    max_age=settings.session_max_age,
)
app.mount("/static", StaticFiles(directory=BASE_DIR / "app" / "static"), name="static")
FRONTEND_DIST = BASE_DIR / "frontend" / "dist"
if FRONTEND_DIST.exists():
    app.mount("/ui", StaticFiles(directory=FRONTEND_DIST, html=True), name="ui")
templates = Jinja2Templates(directory=BASE_DIR / "app" / "templates")
UPLOAD_CHUNK_SIZE = 1024 * 1024
IMPORT_INSERT_CHUNK_SIZE = 1000
IMPORT_PROGRESS_INTERVAL = 5000
analytics_logger = logging.getLogger("analytics_api")
RATE_LIMIT_BUCKETS: dict[str, deque[float]] = defaultdict(deque)
CODING_DAILY_COUNTS: dict[tuple[str, date], int] = defaultdict(int)
CODING_DAILY_ALERTED: set[tuple[str, date]] = set()
coding_logger = logging.getLogger("coding_plan")


class CodingPlanMessage(BaseModel):
    role: Literal["user", "assistant"]
    content: str


class CodingPlanChatPayload(BaseModel):
    messages: list[CodingPlanMessage]


LEGACY_HTML_ROUTES = {
    "/login",
    "/dashboard",
    "/orders",
    "/imports",
    "/analytics/products",
    "/analytics/shops",
    "/settings/users",
    "/settings/roles",
}


@app.middleware("http")
async def prefer_vue_frontend(request: Request, call_next):
    path = request.url.path.rstrip("/") or "/"
    accepts_html = "text/html" in request.headers.get("accept", "")
    legacy_import_detail = path.startswith("/imports/") and request.method == "GET"
    if (
        FRONTEND_DIST.exists()
        and request.method == "GET"
        and accepts_html
        and (path in LEGACY_HTML_ROUTES or legacy_import_detail)
    ):
        return RedirectResponse("/ui/", status_code=303)
    return await call_next(request)


def month_range() -> tuple[str, str]:
    today = date.today()
    start = today.replace(day=1)
    end = add_months(start, 1) - timedelta(days=1)
    return start.isoformat(), end.isoformat()


def add_months(value: date, months: int) -> date:
    month_index = value.month - 1 + months
    year = value.year + month_index // 12
    month = month_index % 12 + 1
    return date(year, month, 1)


def last_month_range() -> tuple[str, str]:
    this_month = date.today().replace(day=1)
    start = add_months(this_month, -1)
    end = this_month - timedelta(days=1)
    return start.isoformat(), end.isoformat()


def render(request: Request, template: str, context: dict[str, Any]) -> HTMLResponse:
    context.setdefault("request", request)
    context.setdefault("current_path", request.url.path)
    return templates.TemplateResponse(template, context)


def redirect(path: str) -> RedirectResponse:
    return RedirectResponse(path, status_code=303)


def current_user(request: Request) -> dict[str, Any] | None:
    user_id = request.session.get("user_id")
    if not user_id:
        return None
    with connection() as conn:
        return load_user_context(conn, int(user_id))


def require_user(request: Request, permission: str | None = None) -> dict[str, Any] | RedirectResponse:
    user = current_user(request)
    if not user:
        return redirect("/login")
    if permission and not has_permission(user, permission):
        raise HTTPException(status_code=403, detail="没有权限访问该功能")
    return user


def require_api_user(request: Request, permission: str | None = None) -> dict[str, Any]:
    user = current_user(request)
    if not user:
        raise HTTPException(status_code=401, detail="未登录")
    if permission and not has_permission(user, permission):
        raise HTTPException(status_code=403, detail="没有权限访问该功能")
    return user


def load_user_by_username(username: str) -> dict[str, Any] | None:
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT id FROM t_user WHERE username = %(username)s AND is_active = 1",
                {"username": username},
            )
            row = cur.fetchone()
            if not row:
                return None
        return load_user_context(conn, int(row["id"]))


def bearer_token(request: Request) -> str:
    header = request.headers.get("authorization", "")
    if header.lower().startswith("bearer "):
        return header[7:].strip()
    return request.headers.get("x-analytics-token", "").strip()


def require_analytics_user(request: Request) -> dict[str, Any]:
    user = current_user(request)
    if user:
        if not has_permission(user, "analytics"):
            raise HTTPException(status_code=403, detail="没有权限访问该功能")
        return user

    token = bearer_token(request)
    if not settings.analytics_api_token or not token or not hmac.compare_digest(token, settings.analytics_api_token):
        raise HTTPException(status_code=401, detail="未登录或缺少有效分析 API Token")

    service_user = load_user_by_username(settings.analytics_api_username)
    if not service_user:
        raise HTTPException(status_code=500, detail="分析 API 服务账号不存在或未启用")
    if not has_permission(service_user, "analytics"):
        raise HTTPException(status_code=403, detail="分析 API 服务账号没有 analytics 权限")
    return service_user


def rate_limit_analytics(request: Request, user: dict[str, Any]) -> None:
    limit = max(1, settings.analytics_rate_limit_per_minute)
    identity = f"{user.get('username') or 'anonymous'}:{request.client.host if request.client else 'unknown'}"
    now = time.monotonic()
    bucket = RATE_LIMIT_BUCKETS[identity]
    while bucket and now - bucket[0] > 60:
        bucket.popleft()
    if len(bucket) >= limit:
        raise HTTPException(status_code=429, detail=f"分析请求过于频繁，请稍后再试（每分钟最多 {limit} 次）")
    bucket.append(now)


def coding_client_ip(request: Request) -> str:
    return request.client.host if request.client else "unknown"


def rate_limit_coding_chat(request: Request, user: dict[str, Any]) -> None:
    limit = max(1, settings.ark_coding_rate_limit_per_minute)
    identity = f"coding:{user.get('username') or 'unknown'}:{coding_client_ip(request)}"
    now = time.monotonic()
    bucket = RATE_LIMIT_BUCKETS[identity]
    while bucket and now - bucket[0] > 60:
        bucket.popleft()
    if len(bucket) >= limit:
        raise HTTPException(status_code=429, detail=f"请求过于频繁，请稍后再试（每分钟最多 {limit} 次）。")
    bucket.append(now)


def consume_coding_daily_quota(user: dict[str, Any]) -> int | None:
    limit = settings.ark_coding_daily_limit_per_user
    if limit <= 0:
        return None

    today = date.today()
    for key in tuple(CODING_DAILY_COUNTS):
        if key[1] < today:
            CODING_DAILY_COUNTS.pop(key, None)
            CODING_DAILY_ALERTED.discard(key)
    username = str(user.get("username") or "unknown")
    key = (username, today)
    next_count = CODING_DAILY_COUNTS[key] + 1
    if next_count > limit:
        raise HTTPException(status_code=429, detail="今日 Coding Plan 调用额度已用完，请明天再试。")

    CODING_DAILY_COUNTS[key] = next_count
    threshold = settings.ark_coding_daily_alert_threshold
    if threshold > 0 and next_count >= threshold and key not in CODING_DAILY_ALERTED:
        CODING_DAILY_ALERTED.add(key)
        coding_logger.warning("coding_daily_alert user=%s calls=%s threshold=%s", username, next_count, threshold)
    return max(0, limit - next_count)


def api_ok(data: Any = None, **extra: Any) -> JSONResponse:
    payload = {"ok": True}
    if data is not None:
        payload["data"] = data
    payload.update(extra)
    return JSONResponse(jsonable_encoder(payload))


def default_filters(request: Request) -> dict[str, Any]:
    start, end = last_month_range()
    return {
        "start_time": request.query_params.get("start_time", start),
        "end_time": request.query_params.get("end_time", end),
        "dept": request.query_params.get("dept", ""),
        "platform": request.query_params.get("platform", ""),
        "shop_name": request.query_params.get("shop_name", ""),
        "category": request.query_params.get("category", ""),
        "product_classification": request.query_params.get("product_classification", ""),
        "product": request.query_params.get("product", ""),
        "order_no": request.query_params.get("order_no", ""),
    }


def order_where(
    user: dict[str, Any],
    filters: dict[str, Any],
    alias: str = "o",
    *,
    validate_dates: bool = True,
) -> tuple[str, dict[str, Any]]:
    if validate_dates:
        validate_date_filters(filters)
    params = dict(filters)
    params["end_time_exclusive"] = (
        datetime.strptime(filters["end_time"], "%Y-%m-%d").date() + timedelta(days=1)
    ).isoformat()
    clauses = [
        f"{alias}.ship_time >= %(start_time)s",
        f"{alias}.ship_time < %(end_time_exclusive)s",
    ]
    if filters.get("category"):
        clauses.append(f"{alias}.category = %(category)s")
    if filters.get("product_classification"):
        clauses.append(f"{alias}.product_classification = %(product_classification)s")
    if filters.get("product"):
        clauses.append(f"({alias}.product_name LIKE %(product_like)s OR {alias}.product_no LIKE %(product_like)s OR {alias}.sku_id LIKE %(product_like)s)")
        params["product_like"] = f"%{filters['product']}%"
    if filters.get("order_no"):
        clauses.append(f"{alias}.order_no LIKE %(order_no_like)s")
        params["order_no_like"] = f"%{filters['order_no']}%"
    scope_filter, params = requested_scope_filters(params, alias)
    where = " WHERE " + " AND ".join(clauses) + scope_filter + scope_clause(user, params, alias)
    return where, params


async def read_json_body(request: Request) -> dict[str, Any]:
    try:
        body = await request.json()
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=400, detail="请求体不是有效 JSON") from exc
    if not isinstance(body, dict):
        raise HTTPException(status_code=400, detail="请求体必须是 JSON 对象")
    return body


def validate_date_filters(filters: dict[str, Any]) -> None:
    try:
        start = datetime.strptime(filters["start_time"], "%Y-%m-%d").date()
        end = datetime.strptime(filters["end_time"], "%Y-%m-%d").date()
    except (KeyError, TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="日期格式必须是 YYYY-MM-DD") from exc
    if start > end:
        raise HTTPException(status_code=400, detail="开始日期不能晚于结束日期")
    if (end - start).days > 366 * 3:
        raise HTTPException(status_code=400, detail="单次分析时间范围不能超过 3 年")


def percentage_change(current: Any, previous: Any) -> Decimal | None:
    current_value = Decimal(current or 0)
    previous_value = Decimal(previous or 0)
    if previous_value == 0:
        return None
    return (current_value - previous_value) / abs(previous_value) * 100


def attach_revenue_shares(rows: list[dict[str, Any]], total_revenue: Any) -> None:
    """Add revenue shares without relying on MySQL 8 window functions."""
    total = Decimal(total_revenue or 0)
    for row in rows:
        revenue = Decimal(row.get("revenue") or 0)
        row["revenue_share_pct"] = revenue / total * 100 if total else Decimal("0")


def attach_exact_product_counts(
    cur,
    rows: list[dict[str, Any]],
    where: str,
    params: dict[str, Any],
    *,
    by_shop: bool = False,
    include_customers: bool = False,
) -> None:
    product_numbers = sorted({str(row.get("product_no") or "") for row in rows if row.get("product_no")})
    if not product_numbers:
        for row in rows:
            row["orders"] = 0
            if include_customers:
                row["customers"] = 0
        return

    count_params = dict(params)
    placeholders: list[str] = []
    for index, product_no in enumerate(product_numbers):
        key = f"candidate_product_{index}"
        count_params[key] = product_no
        placeholders.append(f"%({key})s")
    shop_select = ", o.shop_name" if by_shop else ""
    shop_group = ", o.shop_name" if by_shop else ""
    customer_select = ", COUNT(DISTINCT o.customer_no) AS customers" if include_customers else ""
    cur.execute(
        f"""
        SELECT o.product_no{shop_select},
               COUNT(DISTINCT o.order_no) AS orders
               {customer_select}
        FROM t_order_sku_detail o
        {where}
          AND o.product_no IN ({','.join(placeholders)})
        GROUP BY o.product_no{shop_group}
        """,
        count_params,
    )
    counts = {}
    for count_row in cur.fetchall():
        key = (
            (count_row["product_no"], count_row["shop_name"])
            if by_shop
            else count_row["product_no"]
        )
        counts[key] = count_row
    for row in rows:
        key = (row.get("product_no"), row.get("shop_name")) if by_shop else row.get("product_no")
        count_row = counts.get(key) or {}
        row["orders"] = count_row.get("orders", 0)
        if include_customers:
            row["customers"] = count_row.get("customers", 0)


def mask_sensitive_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    for row in rows:
        if "receiver_name" in row:
            row["receiver_name"] = mask_name(row.get("receiver_name"))
        if "receiver_phone" in row:
            row["receiver_phone"] = mask_phone(row.get("receiver_phone"))
        if "receiver_address" in row:
            row["receiver_address"] = mask_address(row.get("receiver_address"))
    return rows


async def build_analysis_plan(user: dict[str, Any], body: dict[str, Any]) -> tuple[str, dict[str, Any], dict[str, Any], dict[str, Any]]:
    question = str(body.get("question") or "").strip()
    if not question:
        raise HTTPException(status_code=400, detail="请输入要分析的问题")
    if len(question) > 500:
        raise HTTPException(status_code=400, detail="问题长度不能超过 500 字")

    filters = parse_question_filters(question)
    body_filters = body.get("filters") or {}
    if not isinstance(body_filters, dict):
        raise HTTPException(status_code=400, detail="filters 必须是 JSON 对象")
    filters.update({key: value for key, value in body_filters.items() if key in filters and value})
    validate_date_filters(filters)
    where, params = order_where(user, filters)
    ark_intent = await asyncio.to_thread(parse_analytics_intent, question)
    plan = build_sql(question, where, params, intent=ark_intent)
    plan["parser"] = "ark" if ark_intent else "rules"
    plan["ark_enabled"] = ark_analytics_enabled()
    try:
        audit_sql_plan(plan["sql"], params)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=f"分析 SQL 未通过安全审计：{exc}") from exc
    return question, filters, params, plan


def mask_phone(value: str | None) -> str:
    if not value:
        return ""
    text = str(value)
    if len(text) < 7:
        return "***"
    return f"{text[:3]}****{text[-4:]}"


def mask_name(value: str | None) -> str:
    if not value:
        return ""
    text = str(value)
    return text[0] + "**" if len(text) > 1 else "*"


def mask_address(value: str | None) -> str:
    if not value:
        return ""
    text = str(value)
    if len(text) <= 6:
        return "***"
    return f"{text[:3]}***{text[-3:]}"


templates.env.filters["phone_mask"] = mask_phone
templates.env.filters["name_mask"] = mask_name
templates.env.filters["address_mask"] = mask_address


async def save_upload_to_temp(file: UploadFile) -> Path:
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="请上传 xlsx 或 xlsm 文件")

    suffix = Path(file.filename).suffix
    total = 0
    tmp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            tmp_path = Path(tmp.name)
            while True:
                chunk = await file.read(UPLOAD_CHUNK_SIZE)
                if not chunk:
                    break
                total += len(chunk)
                if total > settings.max_upload_bytes:
                    raise HTTPException(status_code=413, detail=f"上传文件不能超过 {settings.max_upload_mb}MB")
                tmp.write(chunk)
        return tmp_path
    except Exception:
        if tmp_path is not None:
            tmp_path.unlink(missing_ok=True)
        raise
    finally:
        await file.close()


@app.get("/healthz")
def healthz():
    return {"ok": True, "env": settings.app_env}


@app.get("/readyz")
def readyz():
    try:
        with connection() as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT 1 AS ok")
                cur.fetchone()
    except Exception as exc:
        raise HTTPException(status_code=503, detail="database unavailable") from exc
    return {"ok": True, "database": "ready"}


@app.get("/", response_class=HTMLResponse)
def index(request: Request):
    if FRONTEND_DIST.exists():
        return redirect("/ui/")
    if current_user(request):
        return redirect("/dashboard")
    return redirect("/login")


@app.get("/favicon.ico")
def favicon():
    return Response(status_code=204)


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    if current_user(request):
        return redirect("/dashboard")
    return render(request, "login.html", {"error": ""})


@app.get("/coding-plan", response_class=HTMLResponse)
def coding_plan_page(request: Request):
    user = require_user(request)
    if isinstance(user, RedirectResponse):
        return user
    return render(request, "coding_plan.html", {})


@app.post("/login")
def login(request: Request, username: str = Form(...), password: str = Form(...)):
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT id, password_hash FROM t_user WHERE username = %(username)s AND is_active = 1", {"username": username})
            user = cur.fetchone()
    if not user or not verify_password(password, user["password_hash"]):
        return render(request, "login.html", {"error": "账号或密码错误"})
    request.session["user_id"] = user["id"]
    return redirect("/dashboard")


@app.post("/logout")
def logout(request: Request):
    request.session.clear()
    return redirect("/login")


@app.get("/api/me")
def api_me(request: Request):
    user = current_user(request)
    if not user:
        return api_ok({"authenticated": False})
    return api_ok(
        {
            "authenticated": True,
            "username": user["username"],
            "display_name": user["display_name"],
            "permissions": sorted(user["permissions"]),
            "role_codes": sorted(user["role_codes"]),
            "scopes": user["scopes"],
            "all_scopes": user["all_scopes"],
        }
    )


@app.post("/api/login")
def api_login(request: Request, username: str = Form(...), password: str = Form(...)):
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT id, password_hash FROM t_user WHERE username = %(username)s AND is_active = 1", {"username": username})
            user = cur.fetchone()
    if not user or not verify_password(password, user["password_hash"]):
        raise HTTPException(status_code=400, detail="账号或密码错误")
    request.session["user_id"] = user["id"]
    return api_me(request)


@app.post("/api/logout")
def api_logout(request: Request):
    request.session.clear()
    return api_ok()


@app.get("/api/coding-plan/status")
def coding_plan_status(request: Request):
    require_api_user(request)
    return api_ok(
        {
            "enabled": coding_chat_enabled(),
            "model": settings.ark_model if coding_chat_enabled() else "未配置",
            "gateway": settings.ark_base_url,
            "max_output_tokens": settings.ark_coding_max_output_tokens,
        }
    )


@app.post("/api/coding-plan/chat")
async def coding_plan_chat(request: Request, payload: CodingPlanChatPayload):
    user = require_api_user(request)
    if not coding_chat_enabled():
        raise HTTPException(status_code=503, detail="服务尚未配置，请联系管理员设置服务端 Ark 环境变量。")
    if not payload.messages:
        raise HTTPException(status_code=400, detail="请先输入问题。")
    if len(payload.messages) > 20:
        raise HTTPException(status_code=400, detail="单次对话最多保留 20 条消息。")

    messages = []
    total_length = 0
    for item in payload.messages:
        content = item.content.strip()
        if not content:
            continue
        if len(content) > 12000:
            raise HTTPException(status_code=400, detail="单条消息不能超过 12000 个字符。")
        total_length += len(content)
        messages.append({"role": item.role, "content": content})
    if not messages:
        raise HTTPException(status_code=400, detail="请先输入问题。")
    if total_length > 40000:
        raise HTTPException(status_code=400, detail="本次对话内容过长，请开启新会话。")

    rate_limit_coding_chat(request, user)
    quota_remaining = consume_coding_daily_quota(user)
    try:
        result = await asyncio.to_thread(ask_coding_plan, messages)
    except ArkCodingError as exc:
        coding_logger.warning(
            "coding_chat_failed user=%s client_ip=%s error=%s",
            user.get("username"),
            coding_client_ip(request),
            type(exc).__name__,
        )
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    usage = result["usage"]
    coding_logger.info(
        "coding_chat_completed user=%s client_ip=%s model=%s prompt_chars=%s usage=%s quota_remaining=%s",
        user.get("username"),
        coding_client_ip(request),
        result["model"],
        total_length,
        usage,
        quota_remaining,
    )
    return api_ok(
        {
            "message": {"role": "assistant", "content": result["content"]},
            "usage": usage,
            "model": result["model"],
            "quota_remaining": quota_remaining,
        }
    )


@app.get("/api/dashboard")
def api_dashboard(request: Request):
    user = require_api_user(request, "view")
    filters = default_filters(request)
    validate_date_filters(filters)
    current_start = datetime.strptime(filters["start_time"], "%Y-%m-%d").date()
    current_end = datetime.strptime(filters["end_time"], "%Y-%m-%d").date()
    period_days = (current_end - current_start).days + 1
    previous_end = current_start - timedelta(days=1)
    previous_start = previous_end - timedelta(days=period_days - 1)
    global_aggregate_available = can_use_global_monthly(user, filters)
    if global_aggregate_available:
        previous_start = add_months(current_start, -1)
        previous_end = current_start - timedelta(days=1)

    combined_filters = dict(filters)
    combined_filters["start_time"] = previous_start.isoformat()
    combined_filters["end_time"] = current_end.isoformat()
    combined_where, combined_params = order_where(user, combined_filters, validate_dates=False)
    combined_params.update(
        {
            "current_start": current_start.isoformat(),
            "current_end_exclusive": (current_end + timedelta(days=1)).isoformat(),
            "previous_start": previous_start.isoformat(),
            "previous_end_exclusive": current_start.isoformat(),
        }
    )
    current_where, current_params = order_where(user, filters)
    trend_format = "%%Y-%%m-%%d" if period_days <= 62 else "%%Y-%%m"
    trend_granularity = "day" if period_days <= 62 else "month"

    with connection() as conn:
        with conn.cursor() as cur:
            if global_aggregate_available:
                monthly_params = {
                    "current_month": current_start.isoformat(),
                    "previous_month": previous_start.isoformat(),
                    "current_end_exclusive": (current_end + timedelta(days=1)).isoformat(),
                }
                cur.execute(
                    """
                    SELECT *
                    FROM agg_dashboard_monthly
                    WHERE stat_month IN (%(current_month)s, %(previous_month)s)
                    """,
                    monthly_params,
                )
                monthly_rows = {row["stat_month"]: row for row in cur.fetchall()}
                current_row = monthly_rows.get(current_start, {})
                previous_row = monthly_rows.get(previous_start, {})
                summary = {
                    "revenue": current_row.get("revenue", Decimal("0")),
                    "qty": current_row.get("qty", Decimal("0")),
                    "profit": current_row.get("profit", Decimal("0")),
                    "orders": current_row.get("orders", 0),
                    "customers": current_row.get("customers", 0),
                    "products": current_row.get("products", 0),
                    "loss_orders": current_row.get("loss_orders", 0),
                    "detail_rows": current_row.get("detail_rows", 0),
                    "latest_ship_time": current_row.get("latest_ship_time"),
                    "previous_revenue": previous_row.get("revenue", Decimal("0")),
                    "previous_profit": previous_row.get("profit", Decimal("0")),
                    "previous_orders": previous_row.get("orders", 0),
                }
                cur.execute(
                    """
                    SELECT DATE_FORMAT(stat_date, '%%Y-%%m-%%d') AS period,
                           revenue, profit, orders
                    FROM agg_dashboard_daily
                    WHERE stat_date >= %(current_month)s
                      AND stat_date < %(current_end_exclusive)s
                    ORDER BY stat_date ASC
                    LIMIT 400
                    """,
                    monthly_params,
                )
                trend_rows = list(cur.fetchall())
                cur.execute(
                    """
                    SELECT platform, revenue, profit, orders
                    FROM agg_platform_monthly
                    WHERE stat_month = %(current_month)s
                    ORDER BY revenue DESC
                    LIMIT 12
                    """,
                    monthly_params,
                )
                platform_rows = list(cur.fetchall())
            else:
                cur.execute(
                    f"""
                    SELECT
                      COALESCE(SUM(CASE WHEN o.ship_time >= %(current_start)s AND o.ship_time < %(current_end_exclusive)s THEN o.share_receivable ELSE 0 END), 0) AS revenue,
                      COALESCE(SUM(CASE WHEN o.ship_time >= %(current_start)s AND o.ship_time < %(current_end_exclusive)s THEN o.qty ELSE 0 END), 0) AS qty,
                      COALESCE(SUM(CASE WHEN o.ship_time >= %(current_start)s AND o.ship_time < %(current_end_exclusive)s THEN o.profit ELSE 0 END), 0) AS profit,
                      COUNT(DISTINCT CASE WHEN o.ship_time >= %(current_start)s AND o.ship_time < %(current_end_exclusive)s THEN o.order_no END) AS orders,
                      COUNT(DISTINCT CASE WHEN o.ship_time >= %(current_start)s AND o.ship_time < %(current_end_exclusive)s THEN o.customer_no END) AS customers,
                      COUNT(DISTINCT CASE WHEN o.ship_time >= %(current_start)s AND o.ship_time < %(current_end_exclusive)s THEN o.product_no END) AS products,
                      COUNT(DISTINCT CASE WHEN o.ship_time >= %(current_start)s AND o.ship_time < %(current_end_exclusive)s AND o.profit < 0 THEN o.order_no END) AS loss_orders,
                      SUM(CASE WHEN o.ship_time >= %(current_start)s AND o.ship_time < %(current_end_exclusive)s THEN 1 ELSE 0 END) AS detail_rows,
                      MAX(CASE WHEN o.ship_time >= %(current_start)s AND o.ship_time < %(current_end_exclusive)s THEN o.ship_time END) AS latest_ship_time,
                      COALESCE(SUM(CASE WHEN o.ship_time >= %(previous_start)s AND o.ship_time < %(previous_end_exclusive)s THEN o.share_receivable ELSE 0 END), 0) AS previous_revenue,
                      COALESCE(SUM(CASE WHEN o.ship_time >= %(previous_start)s AND o.ship_time < %(previous_end_exclusive)s THEN o.profit ELSE 0 END), 0) AS previous_profit,
                      COUNT(DISTINCT CASE WHEN o.ship_time >= %(previous_start)s AND o.ship_time < %(previous_end_exclusive)s THEN o.order_no END) AS previous_orders
                    FROM t_order_sku_detail o
                    {combined_where}
                    """,
                    combined_params,
                )
                summary = cur.fetchone()
                cur.execute(
                    f"""
                    SELECT DATE_FORMAT(o.ship_time, '{trend_format}') AS period,
                           COALESCE(SUM(o.share_receivable), 0) AS revenue,
                           COALESCE(SUM(o.profit), 0) AS profit,
                           COUNT(DISTINCT o.order_no) AS orders
                    FROM t_order_sku_detail o
                    {current_where}
                    GROUP BY DATE_FORMAT(o.ship_time, '{trend_format}')
                    ORDER BY period ASC
                    LIMIT 400
                    """,
                    current_params,
                )
                trend_rows = list(cur.fetchall())
                cur.execute(
                    f"""
                    SELECT o.platform,
                           COALESCE(SUM(o.share_receivable), 0) AS revenue,
                           COALESCE(SUM(o.profit), 0) AS profit,
                           COUNT(DISTINCT o.order_no) AS orders
                    FROM t_order_sku_detail o
                    {current_where}
                    GROUP BY o.platform
                    ORDER BY revenue DESC
                    LIMIT 12
                    """,
                    current_params,
                )
                platform_rows = list(cur.fetchall())
            if can_use_product_monthly(filters):
                aggregate_where, aggregate_params = product_monthly_where(user, filters)
                cur.execute(
                    f"""
                    SELECT a.product_no, MAX(a.product_name) AS product_name,
                           a.category, a.product_classification,
                           SUM(a.qty) AS qty,
                           SUM(a.revenue) AS revenue,
                           SUM(a.profit) AS profit,
                           CASE WHEN SUM(a.revenue) = 0 THEN 0 ELSE SUM(a.profit) / SUM(a.revenue) * 100 END AS profit_rate
                    FROM agg_product_monthly a
                    {aggregate_where}
                    GROUP BY a.product_no, a.category, a.product_classification
                    ORDER BY revenue DESC
                    LIMIT 8
                    """,
                    aggregate_params,
                )
            else:
                cur.execute(
                    f"""
                    SELECT o.product_no, o.product_name, o.category, o.product_classification,
                           SUM(o.qty) AS qty,
                           SUM(o.share_receivable) AS revenue,
                           SUM(o.profit) AS profit,
                           CASE WHEN SUM(o.share_receivable) = 0 THEN 0 ELSE SUM(o.profit) / SUM(o.share_receivable) * 100 END AS profit_rate
                    FROM t_order_sku_detail o
                    {current_where}
                    GROUP BY o.product_no, o.product_name, o.category, o.product_classification
                    ORDER BY revenue DESC
                    LIMIT 8
                    """,
                    current_params,
                )
            top_products = list(cur.fetchall())

    attach_revenue_shares(platform_rows, summary["revenue"])
    attach_revenue_shares(top_products, summary["revenue"])

    revenue = Decimal(summary["revenue"] or 0)
    profit = Decimal(summary["profit"] or 0)
    summary["profit_rate"] = (profit / revenue * 100) if revenue else Decimal("0")
    summary["avg_order_value"] = revenue / summary["orders"] if summary["orders"] else Decimal("0")
    previous_revenue = summary.pop("previous_revenue")
    previous_profit = summary.pop("previous_profit")
    previous_orders = summary.pop("previous_orders")
    comparison = {
        "start_time": previous_start.isoformat(),
        "end_time": previous_end.isoformat(),
        "revenue": previous_revenue,
        "profit": previous_profit,
        "orders": previous_orders,
        "revenue_change_pct": percentage_change(summary["revenue"], previous_revenue),
        "profit_change_pct": percentage_change(summary["profit"], previous_profit),
        "orders_change_pct": percentage_change(summary["orders"], previous_orders),
    }
    return api_ok(
        {
            "filters": filters,
            "summary": summary,
            "comparison": comparison,
            "trend": {"granularity": trend_granularity, "rows": trend_rows},
            "platform_rows": platform_rows,
            "top_products": top_products,
            "meta": {"period_days": period_days},
        }
    )


@app.get("/api/orders")
def api_orders(request: Request):
    user = require_api_user(request, "view")
    filters = default_filters(request)
    where, params = order_where(user, filters)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT o.*
                FROM t_order_sku_detail o
                {where}
                ORDER BY o.ship_time DESC, o.id DESC
                LIMIT 200
                """,
                params,
            )
            rows = list(cur.fetchall())
    return api_ok({"filters": filters, "rows": mask_sensitive_rows(rows)})


@app.get("/api/analytics/products")
def api_products(request: Request):
    user = require_api_user(request, "analytics")
    filters = default_filters(request)
    where, params = order_where(user, filters)
    with connection() as conn:
        with conn.cursor() as cur:
            if can_use_product_monthly(filters):
                aggregate_where, aggregate_params = product_monthly_where(user, filters)
                cur.execute(
                    f"""
                    SELECT a.product_no, MAX(a.product_name) AS product_name,
                           a.category, a.product_classification,
                           SUM(a.qty) AS qty,
                           SUM(a.revenue) AS revenue,
                           SUM(a.cost) AS cost,
                           SUM(a.profit) AS profit,
                           CASE WHEN SUM(a.revenue) = 0 THEN 0 ELSE SUM(a.profit) / SUM(a.revenue) * 100 END AS profit_rate
                    FROM agg_product_monthly a
                    {aggregate_where}
                    GROUP BY a.product_no, a.category, a.product_classification
                    ORDER BY revenue DESC
                    LIMIT 100
                    """,
                    aggregate_params,
                )
                rows = list(cur.fetchall())
                cur.execute(
                    f"""
                    SELECT a.category,
                           SUM(a.qty) AS qty,
                           SUM(a.revenue) AS revenue,
                           SUM(a.cost) AS cost,
                           SUM(a.profit) AS profit,
                           COUNT(DISTINCT a.product_no) AS product_count,
                           CASE WHEN SUM(a.revenue) = 0 THEN 0 ELSE SUM(a.profit) / SUM(a.revenue) * 100 END AS profit_rate
                    FROM agg_product_monthly a
                    {aggregate_where}
                    GROUP BY a.category
                    ORDER BY revenue DESC
                    LIMIT 30
                    """,
                    aggregate_params,
                )
                category_rows = list(cur.fetchall())
            else:
                cur.execute(
                    f"""
                    SELECT o.product_no, o.product_name, o.category, o.product_classification,
                           SUM(o.qty) AS qty,
                           SUM(o.share_receivable) AS revenue,
                           SUM(o.cost) AS cost,
                           SUM(o.profit) AS profit,
                           CASE WHEN SUM(o.share_receivable) = 0 THEN 0 ELSE SUM(o.profit) / SUM(o.share_receivable) * 100 END AS profit_rate
                    FROM t_order_sku_detail o
                    {where}
                    GROUP BY o.product_no, o.product_name, o.category, o.product_classification
                    ORDER BY revenue DESC
                    LIMIT 100
                    """,
                    params,
                )
                rows = list(cur.fetchall())
                cur.execute(
                    f"""
                    SELECT o.category,
                           SUM(o.qty) AS qty,
                           SUM(o.share_receivable) AS revenue,
                           SUM(o.cost) AS cost,
                           SUM(o.profit) AS profit,
                           COUNT(DISTINCT o.product_no) AS product_count,
                           CASE WHEN SUM(o.share_receivable) = 0 THEN 0 ELSE SUM(o.profit) / SUM(o.share_receivable) * 100 END AS profit_rate
                    FROM t_order_sku_detail o
                    {where}
                    GROUP BY o.category
                    ORDER BY revenue DESC
                    LIMIT 30
                    """,
                    params,
                )
                category_rows = list(cur.fetchall())
    return api_ok({"filters": filters, "rows": rows, "category_rows": category_rows})


@app.get("/api/analytics/shops")
def api_shops(request: Request):
    user = require_api_user(request, "analytics")
    filters = default_filters(request)
    where, params = order_where(user, filters)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT o.platform, o.shop_name,
                       SUM(o.qty) AS qty,
                       SUM(o.share_receivable) AS revenue,
                       SUM(o.profit) AS profit,
                       CASE WHEN SUM(o.share_receivable) = 0 THEN 0 ELSE SUM(o.profit) / SUM(o.share_receivable) * 100 END AS profit_rate
                FROM t_order_sku_detail o
                {where}
                GROUP BY o.platform, o.shop_name
                ORDER BY revenue DESC
                LIMIT 100
                """,
                params,
            )
            rows = list(cur.fetchall())
    return api_ok({"filters": filters, "rows": rows})


COMMERCE_DIMENSIONS = {
    "product": {
        "label": "产品",
        "select": "o.product_no AS product_no, o.product_name AS product_name",
        "group": "o.product_no, o.product_name",
        "label_sql": "MAX(o.product_name)",
    },
    "shop": {
        "label": "店铺",
        "select": "o.shop_name AS shop_name",
        "group": "o.shop_name",
        "label_sql": "o.shop_name",
    },
    "platform": {
        "label": "平台",
        "select": "o.platform AS platform",
        "group": "o.platform",
        "label_sql": "o.platform",
    },
    "category": {
        "label": "产品大类",
        "select": "o.category AS category",
        "group": "o.category",
        "label_sql": "o.category",
    },
    "product_classification": {
        "label": "货品分类",
        "select": "o.product_classification AS product_classification",
        "group": "o.product_classification",
        "label_sql": "o.product_classification",
    },
    "province": {
        "label": "省份",
        "select": "o.province AS province",
        "group": "o.province",
        "label_sql": "o.province",
    },
}
COMMERCE_METRICS = {"revenue", "profit", "qty", "orders", "profit_rate"}


@app.get("/api/analytics/commerce-dashboard")
def api_commerce_dashboard(request: Request):
    user = require_api_user(request, "analytics")
    filters = default_filters(request)
    validate_date_filters(filters)
    dimension_key = request.query_params.get("dimension", "product")
    metric = request.query_params.get("metric", "revenue")
    dimension = COMMERCE_DIMENSIONS.get(dimension_key, COMMERCE_DIMENSIONS["product"])
    if metric not in COMMERCE_METRICS:
        metric = "revenue"
    where, params = order_where(user, filters)
    metric_order = "profit_rate" if metric == "profit_rate" else metric
    aggregate_where = ""
    aggregate_params: dict[str, Any] = {}
    aggregate_available = can_use_product_monthly(filters)
    global_aggregate_available = can_use_global_monthly(user, filters)
    if aggregate_available:
        aggregate_where, aggregate_params = product_monthly_where(user, filters)

    with connection() as conn:
        with conn.cursor() as cur:
            if global_aggregate_available:
                summary_params = {"stat_month": filters["start_time"]}
                cur.execute(
                    """
                    SELECT
                      COALESCE(SUM(a.revenue), 0) AS revenue,
                      COALESCE(SUM(a.qty), 0) AS qty,
                      COALESCE(SUM(a.profit), 0) AS profit,
                      COALESCE(SUM(a.orders), 0) AS orders,
                      COALESCE(SUM(a.customers), 0) AS customers,
                      COALESCE(SUM(a.products), 0) AS products,
                      CASE WHEN SUM(a.orders) = 0 THEN 0 ELSE SUM(a.revenue) / SUM(a.orders) END AS avg_order_value,
                      CASE WHEN SUM(a.revenue) = 0 THEN 0 ELSE SUM(a.profit) / SUM(a.revenue) * 100 END AS profit_rate
                    FROM agg_dashboard_monthly a
                    WHERE a.stat_month = %(stat_month)s
                    """,
                    summary_params,
                )
                summary = cur.fetchone()
                cur.execute(
                    """
                    SELECT DATE_FORMAT(a.stat_month, '%%Y-%%m') AS month,
                           a.revenue, a.qty, a.profit, a.orders,
                           CASE WHEN a.revenue = 0 THEN 0 ELSE a.profit / a.revenue * 100 END AS profit_rate
                    FROM agg_dashboard_monthly a
                    WHERE a.stat_month = %(stat_month)s
                    """,
                    summary_params,
                )
                trend_rows = list(cur.fetchall())
            else:
                cur.execute(
                    f"""
                    SELECT
                      COALESCE(SUM(o.share_receivable), 0) AS revenue,
                      COALESCE(SUM(o.qty), 0) AS qty,
                      COALESCE(SUM(o.profit), 0) AS profit,
                      COUNT(DISTINCT o.order_no) AS orders,
                      COUNT(DISTINCT o.customer_no) AS customers,
                      COUNT(DISTINCT o.product_no) AS products,
                      CASE WHEN COUNT(DISTINCT o.order_no) = 0 THEN 0 ELSE SUM(o.share_receivable) / COUNT(DISTINCT o.order_no) END AS avg_order_value,
                      CASE WHEN SUM(o.share_receivable) = 0 THEN 0 ELSE SUM(o.profit) / SUM(o.share_receivable) * 100 END AS profit_rate
                    FROM t_order_sku_detail o
                    {where}
                    """,
                    params,
                )
                summary = cur.fetchone()
                cur.execute(
                    f"""
                    SELECT DATE_FORMAT(o.ship_time, '%%Y-%%m') AS month,
                           COALESCE(SUM(o.share_receivable), 0) AS revenue,
                           COALESCE(SUM(o.qty), 0) AS qty,
                           COALESCE(SUM(o.profit), 0) AS profit,
                           COUNT(DISTINCT o.order_no) AS orders,
                           CASE WHEN SUM(o.share_receivable) = 0 THEN 0 ELSE SUM(o.profit) / SUM(o.share_receivable) * 100 END AS profit_rate
                    FROM t_order_sku_detail o
                    {where}
                    GROUP BY DATE_FORMAT(o.ship_time, '%%Y-%%m')
                    ORDER BY month ASC
                    LIMIT 36
                    """,
                    params,
                )
                trend_rows = list(cur.fetchall())
            if aggregate_available and dimension_key == "product" and metric != "orders":
                cur.execute(
                    f"""
                    SELECT a.product_no, MAX(a.product_name) AS product_name,
                           COALESCE(SUM(a.revenue), 0) AS revenue,
                           COALESCE(SUM(a.qty), 0) AS qty,
                           COALESCE(SUM(a.profit), 0) AS profit,
                           CASE WHEN SUM(a.revenue) = 0 THEN 0 ELSE SUM(a.profit) / SUM(a.revenue) * 100 END AS profit_rate
                    FROM agg_product_monthly a
                    {aggregate_where}
                    GROUP BY a.product_no
                    ORDER BY {metric_order} DESC
                    LIMIT 30
                    """,
                    aggregate_params,
                )
                dimension_rows = list(cur.fetchall())
                attach_exact_product_counts(
                    cur,
                    dimension_rows,
                    where,
                    params,
                    include_customers=True,
                )
            else:
                cur.execute(
                    f"""
                    SELECT {dimension["select"]},
                           COALESCE(SUM(o.share_receivable), 0) AS revenue,
                           COALESCE(SUM(o.qty), 0) AS qty,
                           COALESCE(SUM(o.profit), 0) AS profit,
                           COUNT(DISTINCT o.order_no) AS orders,
                           COUNT(DISTINCT o.customer_no) AS customers,
                           CASE WHEN SUM(o.share_receivable) = 0 THEN 0 ELSE SUM(o.profit) / SUM(o.share_receivable) * 100 END AS profit_rate
                    FROM t_order_sku_detail o
                    {where}
                    GROUP BY {dimension["group"]}
                    ORDER BY {metric_order} DESC
                    LIMIT 30
                    """,
                    params,
                )
                dimension_rows = list(cur.fetchall())

            if aggregate_available:
                cur.execute(
                    f"""
                    SELECT a.product_no, MAX(a.product_name) AS product_name,
                           a.product_classification, a.shop_name,
                           COALESCE(SUM(a.revenue), 0) AS revenue,
                           COALESCE(SUM(a.qty), 0) AS qty,
                           COALESCE(SUM(a.profit), 0) AS profit,
                           CASE WHEN SUM(a.revenue) = 0 THEN 0 ELSE SUM(a.profit) / SUM(a.revenue) * 100 END AS profit_rate
                    FROM agg_product_monthly a
                    {aggregate_where}
                    GROUP BY a.product_no, a.product_classification, a.shop_name
                    HAVING profit < 0 OR profit_rate < 10
                    ORDER BY profit ASC, revenue DESC
                    LIMIT 20
                    """,
                    aggregate_params,
                )
                risk_rows = list(cur.fetchall())
                attach_exact_product_counts(cur, risk_rows, where, params, by_shop=True)
            else:
                cur.execute(
                    f"""
                    SELECT o.product_no, o.product_name, o.product_classification, o.shop_name,
                           COALESCE(SUM(o.share_receivable), 0) AS revenue,
                           COALESCE(SUM(o.qty), 0) AS qty,
                           COALESCE(SUM(o.profit), 0) AS profit,
                           COUNT(DISTINCT o.order_no) AS orders,
                           CASE WHEN SUM(o.share_receivable) = 0 THEN 0 ELSE SUM(o.profit) / SUM(o.share_receivable) * 100 END AS profit_rate
                    FROM t_order_sku_detail o
                    {where}
                    GROUP BY o.product_no, o.product_name, o.product_classification, o.shop_name
                    HAVING profit < 0 OR profit_rate < 10
                    ORDER BY profit ASC, revenue DESC
                    LIMIT 20
                    """,
                    params,
                )
                risk_rows = list(cur.fetchall())

    attach_revenue_shares(dimension_rows, summary["revenue"])

    return api_ok(
        {
            "filters": filters,
            "dimension": {"key": dimension_key if dimension_key in COMMERCE_DIMENSIONS else "product", "label": dimension["label"]},
            "metric": metric,
            "summary": summary,
            "trend_rows": trend_rows,
            "dimension_rows": dimension_rows,
            "risk_rows": risk_rows,
        }
    )


@app.post("/api/analytics/ask")
async def api_analytics_ask(request: Request):
    user = require_analytics_user(request)
    rate_limit_analytics(request, user)
    body = await read_json_body(request)
    question, filters, params, plan = await build_analysis_plan(user, body)
    started = time.perf_counter()
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(plan["sql"], params)
            rows = normalize_rows(list(cur.fetchall()))
    chart = build_chart(question, rows, plan)
    answer = summarize_answer(question, rows, plan)
    analytics_logger.info(
        "analytics ask user=%s rows=%s elapsed_ms=%.1f parser=%s dimensions=%s metrics=%s question=%r",
        user.get("username"),
        len(rows),
        (time.perf_counter() - started) * 1000,
        plan["parser"],
        ",".join(plan["dimensions"]),
        ",".join(plan["metrics"]),
        question,
    )
    return api_ok(
        {
            "question": question,
            "answer": answer,
            "filters": filters,
            "sql": plan["sql"].replace("%%", "%"),
            "sql_params": plan["params"],
            "dimensions": plan["dimensions"],
            "metrics": plan["metrics"],
            "parser": plan["parser"],
            "ark_enabled": plan["ark_enabled"],
            "rows": rows,
            "chart": chart,
        }
    )


@app.post("/api/analytics/parse")
async def api_analytics_parse(request: Request):
    user = require_analytics_user(request)
    rate_limit_analytics(request, user)
    body = await read_json_body(request)
    question, filters, _params, plan = await build_analysis_plan(user, body)
    analytics_logger.info(
        "analytics parse user=%s parser=%s dimensions=%s metrics=%s question=%r",
        user.get("username"),
        plan["parser"],
        ",".join(plan["dimensions"]),
        ",".join(plan["metrics"]),
        question,
    )
    return api_ok(
        {
            "question": question,
            "filters": filters,
            "sql": plan["sql"].replace("%%", "%"),
            "sql_params": plan["params"],
            "dimensions": plan["dimensions"],
            "metrics": plan["metrics"],
            "order_metric": plan["order_metric"],
            "wants_share": plan["wants_share"],
            "parser": plan["parser"],
            "ark_enabled": plan["ark_enabled"],
        }
    )


@app.get("/api/imports")
def api_imports(request: Request):
    require_api_user(request, "import")
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT * FROM t_import_log ORDER BY import_time DESC LIMIT 50")
            logs = list(cur.fetchall())
    return api_ok({"logs": logs})


@app.get("/api/imports/template")
def api_import_template(request: Request):
    require_api_user(request, "import")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "订单导入模板"
    sheet.append(list(HEADER_MAP.keys()))
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 18

    content = BytesIO()
    workbook.save(content)
    content.seek(0)
    return StreamingResponse(
        content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=order_import_template.xlsx"},
    )


@app.post("/api/imports/upload")
async def api_import_upload(request: Request, background_tasks: BackgroundTasks, file: UploadFile = File(...)):
    user = require_api_user(request, "import")
    active_batch = active_import_batch()
    if active_batch:
        raise HTTPException(status_code=409, detail=f"已有导入批次正在处理：{active_batch}，请完成后再上传下一份。")
    batch_no = new_batch_no()
    tmp_path = await save_upload_to_temp(file)
    try:
        create_import_batch(batch_no, file.filename or "", user["username"])
        background_tasks.add_task(process_import_batch, tmp_path, batch_no, file.filename or "", user["username"])
    except Exception:
        tmp_path.unlink(missing_ok=True)
        raise
    return api_ok({"batch_no": batch_no, "status": "processing"})


@app.get("/api/imports/{batch_no}")
def api_import_detail(request: Request, batch_no: str):
    require_api_user(request, "import")
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT * FROM t_import_log WHERE batch_no = %(batch_no)s", {"batch_no": batch_no})
            log = cur.fetchone()
            if not log:
                raise HTTPException(status_code=404, detail="导入批次不存在")
            rows: list[dict[str, Any]] = []
            if log["status"] != "processing":
                cur.execute(
                    """
                    SELECT *
                    FROM tmp_order_import
                    WHERE batch_no = %(batch_no)s
                    ORDER BY COALESCE(error_message, '') <> '' DESC, row_no ASC
                    LIMIT 200
                    """,
                    {"batch_no": batch_no},
                )
                rows = list(cur.fetchall())
    return api_ok({"log": log, "rows": mask_sensitive_rows(rows)})


@app.post("/api/imports/{batch_no}/commit")
def api_import_commit(request: Request, batch_no: str):
    require_api_user(request, "import")
    response = import_commit(request, batch_no)
    if isinstance(response, RedirectResponse):
        return api_ok({"batch_no": batch_no})
    return response


@app.get("/api/settings/users")
def api_settings_users(request: Request):
    require_api_user(request, "settings")
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT u.id, u.username, u.display_name, u.is_active,
                       GROUP_CONCAT(r.role_name ORDER BY r.id SEPARATOR ' / ') AS roles
                FROM t_user u
                LEFT JOIN t_user_role ur ON ur.user_id = u.id
                LEFT JOIN t_role r ON r.id = ur.role_id
                GROUP BY u.id, u.username, u.display_name, u.is_active
                ORDER BY u.id
                """
            )
            users = list(cur.fetchall())
    return api_ok({"users": users})


@app.get("/api/settings/roles")
def api_settings_roles(request: Request):
    require_api_user(request, "settings")
    return api_ok({"roles": load_roles_with_scopes()})


@app.post("/api/settings/roles")
def api_update_role_scope(
    request: Request,
    role_id: int = Form(...),
    dept_scope: str = Form("*"),
    platform_scope: str = Form("*"),
    shop_scope: str = Form("*"),
):
    require_api_user(request, "settings")
    scope_values = {
        "dept": split_scope(dept_scope),
        "platform": split_scope(platform_scope),
        "shop": split_scope(shop_scope),
    }
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM t_role_data_scope WHERE role_id = %(role_id)s", {"role_id": role_id})
            records = []
            for scope_type, values in scope_values.items():
                for value in values:
                    records.append({"role_id": role_id, "scope_type": scope_type, "scope_value": value})
            if records:
                cur.executemany(
                    """
                    INSERT INTO t_role_data_scope (role_id, scope_type, scope_value)
                    VALUES (%(role_id)s, %(scope_type)s, %(scope_value)s)
                    """,
                    records,
                )
    return api_ok({"roles": load_roles_with_scopes()})


@app.get("/dashboard", response_class=HTMLResponse)
def dashboard(request: Request):
    user = require_user(request, "view")
    if isinstance(user, RedirectResponse):
        return user
    filters = default_filters(request)
    where, params = order_where(user, filters)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT
                  COALESCE(SUM(o.share_receivable),0) AS revenue,
                  COALESCE(SUM(o.qty),0) AS qty,
                  COALESCE(SUM(o.profit),0) AS profit,
                  COUNT(DISTINCT o.order_no) AS orders
                FROM t_order_sku_detail o
                {where}
                """,
                params,
            )
            summary = cur.fetchone()
            cur.execute(
                f"""
                SELECT o.product_no, o.product_name, o.product_classification,
                       SUM(o.qty) AS qty,
                       SUM(o.share_receivable) AS revenue,
                       SUM(o.profit) AS profit
                FROM t_order_sku_detail o
                {where}
                GROUP BY o.product_no, o.product_name, o.product_classification
                ORDER BY revenue DESC
                LIMIT 8
                """,
                params,
            )
            top_products = list(cur.fetchall())
    revenue = Decimal(summary["revenue"] or 0)
    profit = Decimal(summary["profit"] or 0)
    summary["profit_rate"] = (profit / revenue * 100) if revenue else Decimal("0")
    return render(request, "dashboard.html", {"user": user, "filters": filters, "summary": summary, "top_products": top_products})


@app.get("/orders", response_class=HTMLResponse)
def orders(request: Request):
    user = require_user(request, "view")
    if isinstance(user, RedirectResponse):
        return user
    filters = default_filters(request)
    where, params = order_where(user, filters)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT o.*
                FROM t_order_sku_detail o
                {where}
                ORDER BY o.ship_time DESC, o.id DESC
                LIMIT 200
                """,
                params,
            )
            rows = list(cur.fetchall())
    return render(request, "orders.html", {"user": user, "filters": filters, "rows": rows})


@app.get("/orders/export")
def orders_export(request: Request):
    user = require_user(request, "export")
    if isinstance(user, RedirectResponse):
        return user
    filters = default_filters(request)
    where, params = order_where(user, filters)
    headers = (
        "order_no", "sku_id", "order_source", "customer_no", "customer_name", "dept", "platform", "shop_name",
        "category", "product_classification", "product_name", "product_no", "unit", "qty", "share_receivable",
        "receiver_name", "receiver_address", "receiver_phone", "province", "city", "district", "ship_time", "cost",
        "express_fee", "logistics_fee", "freight", "aux_material", "share_cost", "profit",
    )

    def safe_csv_value(value: Any) -> Any:
        if value is None:
            return ""
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
            return f"'{value}"
        return value

    def iter_csv():
        import io

        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(headers)
        yield "\ufeff" + buffer.getvalue()
        buffer.seek(0)
        buffer.truncate(0)
        with connection() as conn:
            with conn.cursor(SSDictCursor) as cur:
                cur.execute(
                    f"""
                    SELECT {', '.join(headers)}
                    FROM t_order_sku_detail o
                    {where}
                    ORDER BY o.ship_time DESC
                    LIMIT 50000
                    """,
                    params,
                )
                while True:
                    rows = cur.fetchmany(500)
                    if not rows:
                        break
                    for row in rows:
                        writer.writerow([safe_csv_value(row[column]) for column in headers])
                        yield buffer.getvalue()
                        buffer.seek(0)
                        buffer.truncate(0)

    filename = f"orders_{datetime.now():%Y%m%d%H%M%S}.csv"
    return StreamingResponse(
        iter_csv(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@app.get("/imports", response_class=HTMLResponse)
def imports_page(request: Request):
    user = require_user(request, "import")
    if isinstance(user, RedirectResponse):
        return user
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT * FROM t_import_log ORDER BY import_time DESC LIMIT 50")
            logs = list(cur.fetchall())
    return render(request, "imports.html", {"user": user, "logs": logs, "error": ""})


@app.post("/imports/upload", response_class=HTMLResponse)
async def upload_import(request: Request, background_tasks: BackgroundTasks, file: UploadFile = File(...)):
    user = require_user(request, "import")
    if isinstance(user, RedirectResponse):
        return user
    tmp_path: Path | None = None
    try:
        active_batch = active_import_batch()
        if active_batch:
            raise HTTPException(status_code=409, detail=f"已有导入批次正在处理：{active_batch}，请完成后再上传下一份。")
        batch_no = new_batch_no()
        tmp_path = await save_upload_to_temp(file)
        create_import_batch(batch_no, file.filename or "", user["username"])
        background_tasks.add_task(process_import_batch, tmp_path, batch_no, file.filename or "", user["username"])
    except HTTPException as exc:
        if tmp_path is not None:
            tmp_path.unlink(missing_ok=True)
        return render(request, "imports.html", {"user": user, "logs": [], "error": str(exc.detail)})
    except Exception as exc:
        if tmp_path is not None:
            tmp_path.unlink(missing_ok=True)
        return render(request, "imports.html", {"user": user, "logs": [], "error": str(exc)})
    return redirect(f"/imports/{batch_no}")


IMPORT_COLUMNS = [
    "batch_no",
    "row_no",
    "link_id",
    "sku_id",
    "order_source",
    "customer_no",
    "customer_name",
    "dept",
    "platform",
    "shop_name",
    "order_no",
    "original_order_no",
    "logistics_type",
    "logistics_no",
    "receiver_name",
    "receiver_address",
    "receiver_phone",
    "category",
    "product_classification",
    "product_name",
    "product_no",
    "unit",
    "qty",
    "share_receivable",
    "province",
    "city",
    "district",
    "ship_time",
    "cost",
    "express_fee",
    "logistics_fee",
    "freight",
    "aux_material",
    "share_cost",
    "excel_profit",
    "profit",
    "error_message",
    "warning_message",
]
FINGERPRINT_COLUMNS = tuple(
    column
    for column in IMPORT_COLUMNS
    if column not in {"batch_no", "row_no", "excel_profit", "profit", "error_message", "warning_message"}
)


class UnorderedBatchFingerprint:
    """Order-independent digest for the multiset of rows that will actually be persisted."""

    _MODULUS = 1 << 256

    def __init__(self) -> None:
        self._count = 0
        self._sum = 0

    def update(self, payload: bytes) -> None:
        row_digest = int.from_bytes(hashlib.sha256(payload).digest(), "big")
        self._sum = (self._sum + row_digest) % self._MODULUS
        self._count += 1

    def hexdigest(self) -> str:
        canonical = f"{self._count}:{self._sum:064x}".encode("ascii")
        return hashlib.sha256(canonical).hexdigest()


def insert_tmp_import_rows(cur, rows: list[dict[str, Any]]) -> None:
    placeholders = ",".join([f"%({column})s" for column in IMPORT_COLUMNS])
    cur.executemany(
        f"INSERT INTO tmp_order_import ({','.join(IMPORT_COLUMNS)}) VALUES ({placeholders})",
        rows,
    )


def fingerprint_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return str(value.normalize())
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return value


def update_batch_fingerprint(hasher: Any, row: dict[str, Any]) -> None:
    payload = {column: fingerprint_value(row.get(column)) for column in FINGERPRINT_COLUMNS}
    hasher.update(json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8"))


def create_import_batch(batch_no: str, filename: str, username: str) -> None:
    with connection() as conn:
        with conn.cursor() as cur:
            lock_acquired = False
            try:
                cur.execute("SELECT GET_LOCK('order_manager_import_upload', 5) AS acquired")
                lock_acquired = cur.fetchone()["acquired"] == 1
                if not lock_acquired:
                    raise HTTPException(status_code=409, detail="导入通道繁忙，请稍后重试。")
                cur.execute(
                    """
                    SELECT batch_no
                    FROM t_import_log
                    WHERE status IN ('processing', 'committing')
                    ORDER BY import_time DESC
                    LIMIT 1
                    """
                )
                active = cur.fetchone()
                if active:
                    raise HTTPException(
                        status_code=409,
                        detail=f"已有导入批次正在处理：{active['batch_no']}，请完成后再上传下一份。",
                    )
                cur.execute("DELETE FROM tmp_order_import WHERE batch_no = %(batch_no)s", {"batch_no": batch_no})
                cur.execute(
                    """
                    INSERT INTO t_import_log
                        (batch_no, import_user, file_name, total_rows, success_rows, fail_rows, duplicate_rows, file_hash, status, remark)
                    VALUES
                        (%(batch_no)s, %(user)s, %(file_name)s, 0, 0, 0, 0, '', 'processing', '文件已上传，等待解析')
                    """,
                    {"batch_no": batch_no, "user": username, "file_name": filename},
                )
                conn.commit()
            finally:
                if lock_acquired:
                    cur.execute("SELECT RELEASE_LOCK('order_manager_import_upload')")


def active_import_batch() -> str | None:
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT batch_no
                FROM t_import_log
                WHERE status IN ('processing', 'committing')
                ORDER BY import_time DESC
                LIMIT 1
                """
            )
            row = cur.fetchone()
            return row["batch_no"] if row else None


def update_import_log(batch_no: str, status: str, total_rows: int, fail_rows: int, remark: str) -> None:
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                UPDATE t_import_log
                SET status = %(status)s,
                    total_rows = %(total_rows)s,
                    fail_rows = %(fail_rows)s,
                    import_time = NOW(),
                    remark = %(remark)s
                WHERE batch_no = %(batch_no)s
                """,
                {
                    "batch_no": batch_no,
                    "status": status,
                    "total_rows": total_rows,
                    "fail_rows": fail_rows,
                    "remark": remark[:500],
                },
            )


def finalize_import_batch(batch_no: str, total_rows: int, fail_rows: int, file_hash: str) -> None:
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT batch_no
                FROM t_import_log
                WHERE file_hash = %(file_hash)s
                  AND status = 'committed'
                  AND batch_no <> %(batch_no)s
                LIMIT 1
                """,
                {"batch_no": batch_no, "file_hash": file_hash},
            )
            duplicate_batch = cur.fetchone()
            if duplicate_batch:
                cur.execute("DELETE FROM tmp_order_import WHERE batch_no = %(batch_no)s", {"batch_no": batch_no})
                cur.execute(
                    """
                    UPDATE t_import_log
                    SET total_rows = %(total_rows)s,
                        fail_rows = %(total_rows)s,
                        duplicate_rows = %(total_rows)s,
                        file_hash = %(file_hash)s,
                        status = 'failed',
                        remark = %(remark)s
                    WHERE batch_no = %(batch_no)s
                    """,
                    {
                        "batch_no": batch_no,
                        "total_rows": total_rows,
                        "file_hash": file_hash,
                        "remark": f"整批数据与已入库批次 {duplicate_batch['batch_no']} 重复",
                    },
                )
                return
            cur.execute(
                """
                UPDATE t_import_log l
                SET total_rows = %(total_rows)s,
                    fail_rows = (
                        SELECT COUNT(*) FROM tmp_order_import t
                        WHERE t.batch_no = l.batch_no AND COALESCE(t.error_message, '') <> ''
                    ),
                    duplicate_rows = (
                        SELECT COUNT(*) FROM tmp_order_import t
                        WHERE t.batch_no = l.batch_no AND t.error_message LIKE '%%重复%%'
                    ),
                    file_hash = %(file_hash)s,
                    status = %(status)s,
                    remark = %(remark)s
                WHERE l.batch_no = %(batch_no)s
                """,
                {
                    "batch_no": batch_no,
                    "total_rows": total_rows,
                    "file_hash": file_hash,
                    "status": "failed" if fail_rows else "validated",
                    "remark": (
                        f"校验未通过，已终止：共 {total_rows} 行，异常 {fail_rows} 行"
                        if fail_rows
                        else f"校验通过：共 {total_rows} 行"
                    ),
                },
            )


def process_import_batch(tmp_path: Path, batch_no: str, filename: str, username: str) -> None:
    total_rows = 0
    fail_rows = 0
    buffer: list[dict[str, Any]] = []
    hasher = UnorderedBatchFingerprint()
    try:
        update_import_log(batch_no, "processing", 0, 0, "正在解析 Excel")
        for row in iter_excel_rows(tmp_path, batch_no):
            total_rows += 1
            update_batch_fingerprint(hasher, row)
            if row["error_message"]:
                fail_rows += 1
            buffer.append(row)
            if len(buffer) >= IMPORT_INSERT_CHUNK_SIZE:
                with connection() as conn:
                    with conn.cursor() as cur:
                        insert_tmp_import_rows(cur, buffer)
                buffer.clear()
            if total_rows % IMPORT_PROGRESS_INTERVAL == 0:
                update_import_log(
                    batch_no,
                    "processing",
                    total_rows,
                    fail_rows,
                    f"已解析 {total_rows} 行，异常 {fail_rows} 行",
                )

        if buffer:
            with connection() as conn:
                with conn.cursor() as cur:
                    insert_tmp_import_rows(cur, buffer)

        update_import_log(batch_no, "processing", total_rows, fail_rows, "正在检查整批重复导入")
        finalize_import_batch(batch_no, total_rows, fail_rows, hasher.hexdigest())
    except Exception as exc:
        update_import_log(batch_no, "failed", total_rows, fail_rows, f"导入失败：{exc}")
        try:
            with connection() as conn:
                with conn.cursor() as cur:
                    cur.execute("DELETE FROM tmp_order_import WHERE batch_no = %(batch_no)s", {"batch_no": batch_no})
        except Exception:
            pass
    finally:
        try:
            tmp_path.unlink(missing_ok=True)
        except PermissionError:
            pass
        buffer.clear()
        gc.collect()


def save_import_batch(result: ImportParseResult, filename: str, user: dict[str, Any]) -> None:
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM tmp_order_import WHERE batch_no = %(batch_no)s", {"batch_no": result.batch_no})
            if result.rows:
                insert_tmp_import_rows(cur, result.rows)
            cur.execute(
                """
                INSERT INTO t_import_log
                    (batch_no, import_user, file_name, total_rows, success_rows, fail_rows, duplicate_rows, file_hash, status, remark)
                VALUES
                    (%(batch_no)s, %(user)s, %(file_name)s, %(total_rows)s, 0, %(fail_rows)s, 0, '', %(status)s, %(remark)s)
                ON DUPLICATE KEY UPDATE
                    file_name = VALUES(file_name),
                    total_rows = VALUES(total_rows),
                    fail_rows = VALUES(fail_rows),
                    status = VALUES(status),
                    remark = VALUES(remark)
                """,
                {
                    "batch_no": result.batch_no,
                    "user": user["username"],
                    "file_name": filename,
                    "total_rows": result.total_rows,
                    "fail_rows": result.fail_rows,
                    "status": "failed" if result.fail_rows else "validated",
                    "remark": (
                        f"校验未通过，已终止：{result.error_summary}"
                        if result.fail_rows
                        else f"校验通过：{result.error_summary}"
                    ),
                },
            )
            cur.execute(
                """
                UPDATE t_import_log l
                SET fail_rows = (
                    SELECT COUNT(*) FROM tmp_order_import t
                    WHERE t.batch_no = l.batch_no AND COALESCE(t.error_message, '') <> ''
                ),
                duplicate_rows = (
                    SELECT COUNT(*) FROM tmp_order_import t
                    WHERE t.batch_no = l.batch_no AND t.error_message LIKE '%%重复%%'
                )
                WHERE l.batch_no = %(batch_no)s
                """,
                {"batch_no": result.batch_no},
            )


@app.get("/imports/{batch_no}", response_class=HTMLResponse)
def import_detail(request: Request, batch_no: str):
    user = require_user(request, "import")
    if isinstance(user, RedirectResponse):
        return user
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT * FROM t_import_log WHERE batch_no = %(batch_no)s", {"batch_no": batch_no})
            log = cur.fetchone()
            if not log:
                raise HTTPException(status_code=404, detail="导入批次不存在")
            cur.execute(
                """
                SELECT *
                FROM tmp_order_import
                WHERE batch_no = %(batch_no)s
                ORDER BY COALESCE(error_message, '') <> '' DESC, row_no ASC
                LIMIT 200
                """,
                {"batch_no": batch_no},
            )
            rows = list(cur.fetchall())
    return render(request, "import_detail.html", {"user": user, "log": log, "rows": rows})


@app.post("/imports/{batch_no}/commit")
def import_commit(request: Request, batch_no: str):
    user = require_user(request, "import")
    if isinstance(user, RedirectResponse):
        return user
    with connection() as conn:
        lock_name = ""
        lock_acquired = False
        try:
            with conn.cursor() as cur:
                cur.execute("SELECT status, file_hash FROM t_import_log WHERE batch_no = %(batch_no)s", {"batch_no": batch_no})
                log = cur.fetchone()
                if not log:
                    raise HTTPException(status_code=404, detail="导入批次不存在")

                lock_seed = str(log.get("file_hash") or batch_no)
                lock_name = f"order_commit_{hashlib.sha256(lock_seed.encode('utf-8')).hexdigest()[:40]}"
                cur.execute("SELECT GET_LOCK(%(lock_name)s, 10) AS acquired", {"lock_name": lock_name})
                lock_acquired = cur.fetchone()["acquired"] == 1
                if not lock_acquired:
                    raise HTTPException(status_code=409, detail="批次正在提交，请稍后查看结果。")

                cur.execute("SELECT status, file_hash FROM t_import_log WHERE batch_no = %(batch_no)s", {"batch_no": batch_no})
                log = cur.fetchone()
                if log["status"] == "committed":
                    return redirect(f"/imports/{batch_no}")
                if log["status"] != "validated":
                    raise HTTPException(status_code=400, detail="批次尚未校验完成，不能确认入库")

                if log.get("file_hash"):
                    cur.execute(
                        """
                        SELECT batch_no
                        FROM t_import_log
                        WHERE file_hash = %(file_hash)s
                          AND status = 'committed'
                          AND batch_no <> %(batch_no)s
                        LIMIT 1
                        """,
                        {"file_hash": log["file_hash"], "batch_no": batch_no},
                    )
                    duplicate = cur.fetchone()
                    if duplicate:
                        cur.execute("DELETE FROM tmp_order_import WHERE batch_no = %(batch_no)s", {"batch_no": batch_no})
                        cur.execute(
                            """
                            UPDATE t_import_log
                            SET status = 'failed', fail_rows = total_rows, duplicate_rows = total_rows,
                                remark = %(remark)s, import_time = NOW()
                            WHERE batch_no = %(batch_no)s
                            """,
                            {
                                "batch_no": batch_no,
                                "remark": f"整批数据与已入库批次 {duplicate['batch_no']} 重复",
                            },
                        )
                        conn.commit()
                        return redirect(f"/imports/{batch_no}")

                cur.execute(
                    """
                    UPDATE t_import_log
                    SET status = 'committing', remark = '正在写入正式订单'
                    WHERE batch_no = %(batch_no)s AND status = 'validated'
                    """,
                    {"batch_no": batch_no},
                )
                if cur.rowcount != 1:
                    raise HTTPException(status_code=409, detail="批次状态已变化，请刷新后重试。")
                cur.execute(
                    "SELECT COUNT(*) AS c FROM tmp_order_import WHERE batch_no = %(batch_no)s AND COALESCE(error_message, '') <> ''",
                    {"batch_no": batch_no},
                )
                if cur.fetchone()["c"]:
                    raise HTTPException(status_code=400, detail="存在异常数据，不能确认入库")
                cur.execute(
                    """
                    SELECT DISTINCT DATE_FORMAT(ship_time, '%%Y-%%m-01') AS stat_month
                    FROM tmp_order_import
                    WHERE batch_no = %(batch_no)s
                    ORDER BY stat_month
                    """,
                    {"batch_no": batch_no},
                )
                affected_months = [row["stat_month"] for row in cur.fetchall()]
                cur.execute(
                    """
                    INSERT INTO t_order_sku_detail (
                      link_id, sku_id, order_source, customer_no, customer_name, dept, platform, shop_name, order_no, original_order_no,
                      logistics_type, logistics_no, receiver_name, receiver_address, receiver_phone, category, product_classification, product_name,
                      product_no, unit, qty, share_receivable, province, city, district, ship_time, cost, express_fee,
                      logistics_fee, freight, aux_material, share_cost
                    )
                    SELECT
                      link_id, sku_id, order_source, customer_no, customer_name, dept, platform, shop_name, order_no, original_order_no,
                      logistics_type, logistics_no, receiver_name, receiver_address, receiver_phone, category, product_classification, product_name,
                      product_no, unit, qty, share_receivable, province, city, district, ship_time, cost, express_fee,
                      logistics_fee, freight, aux_material, share_cost
                    FROM tmp_order_import
                    WHERE batch_no = %(batch_no)s
                    """,
                    {"batch_no": batch_no},
                )
                success_rows = cur.rowcount
                refresh_product_months(conn, affected_months)
                refresh_dashboard_months(conn, affected_months)
                cur.execute(
                    """
                    UPDATE t_import_log
                    SET success_rows = %(success_rows)s, fail_rows = 0, status = 'committed',
                        remark = %(remark)s, import_time = NOW()
                    WHERE batch_no = %(batch_no)s AND status = 'committing'
                    """,
                    {
                        "success_rows": success_rows,
                        "batch_no": batch_no,
                        "remark": f"已入库：共 {success_rows} 行",
                    },
                )
                cur.execute("DELETE FROM tmp_order_import WHERE batch_no = %(batch_no)s", {"batch_no": batch_no})
                conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            if lock_acquired:
                with conn.cursor() as cur:
                    cur.execute("SELECT RELEASE_LOCK(%(lock_name)s)", {"lock_name": lock_name})
    return redirect(f"/imports/{batch_no}")


@app.get("/analytics/products", response_class=HTMLResponse)
def product_analytics(request: Request):
    user = require_user(request, "analytics")
    if isinstance(user, RedirectResponse):
        return user
    filters = default_filters(request)
    where, params = order_where(user, filters)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT o.product_no, o.product_name, o.category, o.product_classification,
                       SUM(o.qty) AS qty,
                       SUM(o.share_receivable) AS revenue,
                       SUM(o.cost) AS cost,
                       SUM(o.profit) AS profit,
                       CASE WHEN SUM(o.share_receivable) = 0 THEN 0 ELSE SUM(o.profit) / SUM(o.share_receivable) * 100 END AS profit_rate
                FROM t_order_sku_detail o
                {where}
                GROUP BY o.product_no, o.product_name, o.category, o.product_classification
                ORDER BY revenue DESC
                LIMIT 100
                """,
                params,
            )
            rows = list(cur.fetchall())
            cur.execute(
                f"""
                SELECT o.category,
                       SUM(o.qty) AS qty,
                       SUM(o.share_receivable) AS revenue,
                       SUM(o.cost) AS cost,
                       SUM(o.profit) AS profit,
                       COUNT(DISTINCT o.product_no) AS product_count,
                       CASE WHEN SUM(o.share_receivable) = 0 THEN 0 ELSE SUM(o.profit) / SUM(o.share_receivable) * 100 END AS profit_rate
                FROM t_order_sku_detail o
                {where}
                GROUP BY o.category
                ORDER BY revenue DESC
                LIMIT 30
                """,
                params,
            )
            category_rows = list(cur.fetchall())
    max_product_revenue = max([float(row["revenue"] or 0) for row in rows] or [0])
    max_category_revenue = max([float(row["revenue"] or 0) for row in category_rows] or [0])
    max_category_profit = max([abs(float(row["profit"] or 0)) for row in category_rows] or [0])
    for row in rows:
        revenue = float(row["revenue"] or 0)
        profit = float(row["profit"] or 0)
        row["revenue_pct"] = 0 if max_product_revenue == 0 else revenue / max_product_revenue * 100
        row["profit_pct"] = 0 if revenue == 0 else max(min(profit / revenue * 100, 100), -100)
    for row in category_rows:
        revenue = float(row["revenue"] or 0)
        profit = float(row["profit"] or 0)
        row["revenue_pct"] = 0 if max_category_revenue == 0 else revenue / max_category_revenue * 100
        row["profit_pct"] = 0 if max_category_profit == 0 else abs(profit) / max_category_profit * 100
        row["profit_width"] = row["profit_pct"]
        row["profit_negative"] = profit < 0
    return render(
        request,
        "analytics_products.html",
        {"user": user, "filters": filters, "rows": rows, "category_rows": category_rows},
    )


@app.get("/analytics/shops", response_class=HTMLResponse)
def shop_analytics(request: Request):
    user = require_user(request, "analytics")
    if isinstance(user, RedirectResponse):
        return user
    filters = default_filters(request)
    where, params = order_where(user, filters)
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT o.platform, o.shop_name,
                       SUM(o.qty) AS qty,
                       SUM(o.share_receivable) AS revenue,
                       SUM(o.profit) AS profit,
                       CASE WHEN SUM(o.share_receivable) = 0 THEN 0 ELSE SUM(o.profit) / SUM(o.share_receivable) * 100 END AS profit_rate
                FROM t_order_sku_detail o
                {where}
                GROUP BY o.platform, o.shop_name
                ORDER BY revenue DESC
                LIMIT 100
                """,
                params,
            )
            rows = list(cur.fetchall())
    return render(request, "analytics_shops.html", {"user": user, "filters": filters, "rows": rows})


@app.get("/settings/users", response_class=HTMLResponse)
def settings_users(request: Request):
    user = require_user(request, "settings")
    if isinstance(user, RedirectResponse):
        return user
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT u.id, u.username, u.display_name, u.is_active,
                       GROUP_CONCAT(r.role_name ORDER BY r.id SEPARATOR ' / ') AS roles
                FROM t_user u
                LEFT JOIN t_user_role ur ON ur.user_id = u.id
                LEFT JOIN t_role r ON r.id = ur.role_id
                GROUP BY u.id, u.username, u.display_name, u.is_active
                ORDER BY u.id
                """
            )
            users = list(cur.fetchall())
    return render(request, "settings_users.html", {"user": user, "users": users})


@app.get("/settings/roles", response_class=HTMLResponse)
def settings_roles(request: Request):
    user = require_user(request, "settings")
    if isinstance(user, RedirectResponse):
        return user
    roles = load_roles_with_scopes()
    return render(request, "settings_roles.html", {"user": user, "roles": roles, "message": ""})


@app.post("/settings/roles", response_class=HTMLResponse)
def update_role_scope(
    request: Request,
    role_id: int = Form(...),
    dept_scope: str = Form("*"),
    platform_scope: str = Form("*"),
    shop_scope: str = Form("*"),
):
    user = require_user(request, "settings")
    if isinstance(user, RedirectResponse):
        return user
    scope_values = {
        "dept": split_scope(dept_scope),
        "platform": split_scope(platform_scope),
        "shop": split_scope(shop_scope),
    }
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM t_role_data_scope WHERE role_id = %(role_id)s", {"role_id": role_id})
            records = []
            for scope_type, values in scope_values.items():
                for value in values:
                    records.append({"role_id": role_id, "scope_type": scope_type, "scope_value": value})
            if records:
                cur.executemany(
                    """
                    INSERT INTO t_role_data_scope (role_id, scope_type, scope_value)
                    VALUES (%(role_id)s, %(scope_type)s, %(scope_value)s)
                    """,
                    records,
                )
    return render(request, "settings_roles.html", {"user": user, "roles": load_roles_with_scopes(), "message": "角色数据范围已保存"})


def split_scope(value: str) -> list[str]:
    cleaned = [part.strip() for part in value.replace("，", ",").split(",") if part.strip()]
    if not cleaned:
        return []
    if "*" in cleaned:
        return ["*"]
    return cleaned


def load_roles_with_scopes() -> list[dict[str, Any]]:
    with connection() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT * FROM t_role ORDER BY id")
            roles = list(cur.fetchall())
            cur.execute("SELECT role_id, scope_type, scope_value FROM t_role_data_scope ORDER BY scope_type, scope_value")
            rows = list(cur.fetchall())
    for role in roles:
        role["scopes"] = {"dept": [], "platform": [], "shop": []}
        for row in rows:
            if row["role_id"] == role["id"]:
                role["scopes"][row["scope_type"]].append(row["scope_value"])
        for key, values in role["scopes"].items():
            role["scopes"][key] = ",".join(values)
    return roles


@app.exception_handler(403)
def forbidden(request: Request, exc: HTTPException):
    if request.url.path.startswith("/api/"):
        return JSONResponse({"detail": exc.detail}, status_code=403)
    user = current_user(request)
    return templates.TemplateResponse(
        "error.html",
        {"request": request, "current_path": request.url.path, "user": user, "title": "没有权限", "message": exc.detail},
        status_code=403,
    )
