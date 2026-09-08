from __future__ import annotations

import io
import json
import re
import unittest
from decimal import Decimal

from openpyxl import load_workbook

from app.report_exports import (
    _pdf_number,
    _pdf_operating_values,
    _report_metadata_rows,
    export_exceptions_xlsx,
    export_report_pdf,
    export_report_xlsx,
)


def exception(index: int, *, quality: bool = False) -> dict:
    formula = '=HYPERLINK("https://example.invalid","x")'
    return {
        "id": f"row-{index}",
        "type": "quality" if quality else "operating",
        "rule_id": "data_error" if quality else "negative_profit",
        "rule_label": "数据错误" if quality else "负利润",
        "severity": "high",
        "object_label": formula if index == 0 else f"商品{index}",
        "platform": "平台A",
        "shop_name": "门店A",
        "product_no": f"P{index}",
        "current_value": Decimal("-10.25"),
        "baseline_value": Decimal("1.25"),
        "change_pct": Decimal("-920"),
        "impact_amount": Decimal("11.50"),
        "message": "需要核查",
        "suggestion": "核对来源数据",
        "status": "confirmed",
        "note": "已核对渠道结算单，备注内容需要随 Top N 输出。",
        "reviewed_at": "2026-09-08T16:00:00",
        "evidence_hash": f"hash-{index}",
        "evidence": {"batch_no": "B001", "row_no": index + 2, "error_message": formula, "warning_message": ""},
    }


def report() -> dict:
    operating = [exception(i) for i in range(65)]
    quality = [exception(i, quality=True) for i in range(3)]
    return {
        "id": "report-001",
        "version": 3,
        "created_at": "2026-09-08T16:10:00",
        "title": "六月经营分析报告",
        "report_type": "monthly",
        "filters": {
            "start_time": "2026-06-01", "end_time": "2026-06-30",
            "dept": "华东", "platform": "平台A", "shop_name": "门店A",
            "category": "食品", "product_classification": "常温",
            "product": "P1", "province": "上海", "city": "上海市", "order_source": "直播",
        },
        "comparison": {"label": "2026-05-01 至 2026-05-31", "mode": "previous_period"},
        "summary": {"revenue": Decimal("5500"), "profit": Decimal("380"), "profit_rate": Decimal("6.909"), "qty": Decimal("13"), "orders": 4, "detail_rows": 5, "fees": Decimal("650")},
        "previous_summary": {"revenue": Decimal("4200"), "profit": Decimal("700")},
        "contributions": [{"object_label": "商品P5", "revenue_change": Decimal("-1000"), "profit_change": Decimal("-300")}],
        "operating_rows": operating,
        "quality_rows": quality,
        "rules": {"low_margin_pct": 5, "high_fee_pct": 20, "decline_pct": 20, "min_revenue": 1000},
        "rules_version": "rules-v1",
        "data_version": "data-v1",
        "warnings": ["扣减费用率仅含运费、辅料及分摊费用。"],
        "narrative": ["变化贡献仅表示数值差额，不代表已验证的业务原因。"],
        "sections": [{"key": "operating", "title": "重点经营异常（Top 10）", "top_n": 10, "rows": operating[:10]}],
    }


class ReportExportTests(unittest.TestCase):
    def test_report_context_contains_identity_scope_versions_and_rules(self) -> None:
        rows = dict(_report_metadata_rows(report()))
        self.assertEqual(rows["报告ID"], "report-001")
        self.assertEqual(rows["版本"], 3)
        self.assertEqual(rows["生成时间"], "2026-09-08T16:10:00")
        self.assertEqual(rows["部门"], "华东")
        self.assertEqual(rows["平台"], "平台A")
        self.assertEqual(rows["店铺"], "门店A")
        self.assertEqual(rows["产品大类"], "食品")
        self.assertEqual(rows["货品分类"], "常温")
        self.assertEqual(rows["低利润率阈值"], "5%")
        self.assertEqual(rows["最低销售额"], 1000)
        self.assertEqual(rows["规则版本"], "rules-v1")
        self.assertEqual(rows["数据版本"], "data-v1")

    def test_pdf_top_operating_values_include_review_status_and_note(self) -> None:
        values = _pdf_operating_values(exception(1))
        self.assertEqual(values["status"], "confirmed")
        self.assertIn("渠道结算单", values["note"])

    def test_pdf_formats_archived_numeric_strings_to_two_decimals(self) -> None:
        self.assertEqual(_pdf_number("6.9090909090909090909"), "6.91")
        self.assertEqual(_pdf_number("not-a-number"), "not-a-number")

    def test_exception_workbook_contains_every_row_and_escapes_formula_strings(self) -> None:
        rows = [exception(i) for i in range(205)]
        payload = export_exceptions_xlsx(rows)
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 206)
        headers = [cell.value for cell in sheet[1]]
        object_col = headers.index("对象") + 1
        evidence_col = headers.index("证据") + 1
        self.assertTrue(sheet.cell(2, object_col).value.startswith("'="))
        self.assertNotEqual(sheet.cell(2, object_col).data_type, "f")
        self.assertNotIn('"error_message": "=', sheet.cell(2, evidence_col).value)

    def test_report_workbook_keeps_full_exception_sheets_and_numeric_summary(self) -> None:
        archived_report = json.loads(json.dumps(report(), ensure_ascii=False, default=str))
        payload = export_report_xlsx(archived_report)
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        self.assertEqual(workbook["经营异常"].max_row, 66)
        self.assertEqual(workbook["数据质量异常"].max_row, 4)
        summary = workbook["经营摘要"]
        values = {summary.cell(row, 1).value: summary.cell(row, 2).value for row in range(2, summary.max_row + 1)}
        self.assertEqual(values["销售额"], 5500)
        self.assertEqual(values["利润"], 380)
        self.assertEqual(values["报告ID"], "report-001")
        self.assertEqual(values["版本"], 3)
        self.assertEqual(values["生成时间"], "2026-09-08T16:10:00")
        self.assertEqual(summary.cell(2, 2).data_type, "n")
        contribution_headers = [cell.value for cell in workbook["变化贡献"][1]]
        revenue_change_col = contribution_headers.index("销售额贡献") + 1
        self.assertEqual(workbook["变化贡献"].cell(2, revenue_change_col).data_type, "n")
        exception_headers = [cell.value for cell in workbook["经营异常"][1]]
        impact_col = exception_headers.index("规则影响额") + 1
        self.assertEqual(workbook["经营异常"].cell(2, impact_col).data_type, "n")

    def test_exception_workbook_includes_review_fields(self) -> None:
        row = {**exception(1), "status": "confirmed", "note": "已核对", "reviewed_at": "2026-09-08T16:00:00"}
        workbook = load_workbook(io.BytesIO(export_exceptions_xlsx([row])))
        headers = [cell.value for cell in workbook.active[1]]
        self.assertIn("核查状态", headers)
        self.assertIn("核查备注", headers)
        self.assertIn("核查时间", headers)

    def test_pdf_has_chinese_report_content_top_n_label_and_multiple_pages(self) -> None:
        payload = export_report_pdf(report())
        self.assertTrue(payload.startswith(b"%PDF-"))
        self.assertGreaterEqual(len(re.findall(rb"/Type\s*/Page\b", payload)), 2)
        self.assertIn(b"Top 10", payload)
        self.assertGreater(len(payload), 5000)


if __name__ == "__main__":
    unittest.main()
